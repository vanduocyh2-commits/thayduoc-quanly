"""Teacher-only offline exam pack: editable Word equations, Excel keys, and QR metadata."""
import io,json,re,zipfile,html
from collections import defaultdict
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from latex2mathml.converter import convert as latex_mathml
from mathml2omml import convert as mathml_omml
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
import qrcode
import qrcode.image.svg

TYPES={'quiz':'Trắc nghiệm 4 lựa chọn','truefalse':'Đúng / Sai (4 ý)','short':'Trả lời ngắn','essay':'Tự luận'}
LEVELS=['Nhận biết','Thông hiểu','Vận dụng','Vận dụng cao']
MATH=re.compile(r'(\$\$[\s\S]+?\$\$|\$(?!\$)[^\n$]+?\$|\\\([\s\S]+?\\\)|\\\[[\s\S]+?\\\])')
def formula_value(token):return token[2:-2] if token.startswith(('$$','\\(','\\[')) else token[1:-1]
def add_content(p,value):
 for piece in MATH.split(str(value)):
  if not piece:continue
  if MATH.fullmatch(piece):
   try:
    eq=mathml_omml(latex_mathml(formula_value(piece)))
    root=parse_xml(('<root xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'+eq+'</root>').encode())
    p._p.append(root[0])
   except Exception:p.add_run(piece)
  else:p.add_run(piece.replace('**','').replace('`',''))
def new_doc(title):
 doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(1.8)
 doc.styles['Normal'].font.name='Times New Roman';doc.styles['Normal'].font.size=Pt(12)
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(title).bold=True
 return doc
def doc_bytes(doc):
 out=io.BytesIO();doc.save(out);return out.getvalue()
def key(q):
 kind=q.get('type','quiz')
 if kind=='essay':return 'Giáo viên chấm theo bài làm'
 if kind=='truefalse':return '; '.join('abcd'[j]+'. '+('Đúng' if a else 'Sai') for j,a in enumerate(q['correct']))
 if kind=='short':return q['correct']
 return 'ABCD'[q['correct']]
def exam_doc(data,variant):
 doc=new_doc(data['title']);doc.add_paragraph(f"Môn: {data.get('subject','')} · Khối: {data.get('grade','')} · Thời gian: {data.get('minutes',45)} phút · Mã đề: {variant['code']}")
 doc.add_paragraph('Họ tên: .......................................................... Lớp: ................')
 add_content(doc.add_paragraph(),data.get('body',''))
 qs=variant['questions']+variant['essays'];last_kind=None
 for i,q in enumerate(qs):
  kind=q.get('type','quiz')
  if kind!=last_kind:doc.add_heading(TYPES[kind],level=2);last_kind=kind
  p=doc.add_paragraph();p.add_run(f"Câu {i+1}. ({q.get('points','')} điểm) ").bold=True;add_content(p,q['text'])
  if kind in ['quiz','truefalse']:
   for j,a in enumerate(q['answers']):
    p=doc.add_paragraph();p.add_run(('ABCD' if kind=='quiz' else 'abcd')[j]+'. ');add_content(p,a)
  else:doc.add_paragraph('........................................................................................')
 doc.add_paragraph('— HẾT —').alignment=WD_ALIGN_PARAGRAPH.CENTER
 return doc_bytes(doc)
def matrix_doc(data,variant):
 doc=new_doc('MA TRẬN & BẢN ĐẶC TẢ — DÀNH CHO GIÁO VIÊN');doc.add_paragraph(data['title']+' · Mã đề '+variant['code']);qs=variant['questions']+variant['essays']
 topics=list(dict.fromkeys(q['topic'] for q in qs));table=doc.add_table(rows=1,cols=6);table.style='Table Grid'
 for cell,value in zip(table.rows[0].cells,['Chủ đề']+LEVELS+['Tổng điểm']):cell.text=value
 for topic in topics:
  a=[q for q in qs if q['topic']==topic];cells=table.add_row().cells;cells[0].text=topic
  for i,level in enumerate(LEVELS):
   group=[q for q in a if q['level']==level];cells[i+1].text=f"{len(group)} câu / {sum(q['points'] for q in group):g} đ"
  cells[5].text=f"{sum(q['points'] for q in a):g}"
 doc.add_heading('Đặc tả & đáp án',level=2)
 for i,q in enumerate(qs):
  doc.add_heading(f"Câu {i+1} · {TYPES[q.get('type','quiz')]} · {q['level']} · {q['points']:g} điểm",level=3)
  add_content(doc.add_paragraph(),q.get('outcome',''));add_content(doc.add_paragraph(),q['text']);add_content(doc.add_paragraph(),key(q))
 return doc_bytes(doc)
def outline_doc(data):
 doc=new_doc('ĐỀ CƯƠNG ÔN TẬP');doc.add_paragraph(data['title']+' · '+data.get('subject',''));topics=defaultdict(list)
 for q in data['questions']+data.get('essays',[]):
  value=q.get('outcome') or q['text']
  if value not in topics[q['topic']]:topics[q['topic']].append(value)
 for topic,outcomes in topics.items():
  doc.add_heading(topic,level=2)
  for outcome in outcomes:add_content(doc.add_paragraph(style='List Bullet'),outcome)
 return doc_bytes(doc)
def export_pack(data):
 out=io.BytesIO();wb=Workbook();ws=wb.active;ws.title='Đáp án';ws.append(['Mã đề','Thứ tự câu','Dạng','Chủ đề','Mức độ','Điểm','Đáp án','Yêu cầu cần đạt'])
 variants=data.get('variants') or [dict(code='GOC',questions=data['questions'],essays=data.get('essays',[]))]
 with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
  for v in variants:
   z.writestr('de-'+v['code']+'.docx',exam_doc(data,v))
   if data.get('includeMatrix'):z.writestr('ma-tran-dac-ta-'+v['code']+'.docx',matrix_doc(data,v))
   for i,q in enumerate(v['questions']+v['essays']):
    values=[v['code'],i+1,TYPES[q.get('type','quiz')],q['topic'],q['level'],q['points'],key(q),q.get('outcome','')];ws.append(values)
    for cell in ws[ws.max_row]:
     if isinstance(cell.value,str):cell.data_type='s' # Never execute teacher text as an Excel formula.
   image=qrcode.make(json.dumps({'schema':'school-exam-v1','subjectCode':data.get('subjectCode','00'),'examCode':v['code'],'questions':len(v['questions'])+len(v['essays'])}),image_factory=qrcode.image.svg.SvgPathImage);qr=io.BytesIO();image.save(qr);z.writestr('qr-ma-de-'+v['code']+'.svg',qr.getvalue())
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
  for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='237B65')
  for col,width in zip('ABCDEFGH',[12,14,28,30,20,12,42,60]):ws.column_dimensions[col].width=width
  for row in ws:
   for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
  book=io.BytesIO();wb.save(book);z.writestr('dap-an.xlsx',book.getvalue());z.writestr('de-cuong.docx',outline_doc(data))
  z.writestr('huong-dan.txt','Gói dành cho giáo viên. Không gửi đáp án Excel hoặc bản đặc tả cho học sinh.\nĐề Word dùng công thức Equation của Word khi chuyển đổi được; công thức không hỗ trợ giữ nguyên LaTeX để kiểm tra.\nQR chứa mã môn, mã đề và số câu; không chứa đáp án, không tự kết nối ứng dụng chấm thi khác.\nCâu được đánh số liên tục để đối chiếu đúng giữa đề Word, bài làm trên web và đáp án Excel (phần tự chấm trước, tự luận sau).\n')
 return out.getvalue()
