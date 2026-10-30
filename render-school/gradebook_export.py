"""Class-specific, formatted native Excel gradebooks."""
import io,re,unicodedata
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
from exam_results import with_best_result,valid_score

def ascii_name(value):
    text=''.join(c for c in unicodedata.normalize('NFD',value.replace('đ','d').replace('Đ','D')) if not unicodedata.combining(c))
    return re.sub(r'[^A-Z0-9]+','_',text.upper()).strip('_')
def day(value):
    if not value:return ''
    try:return datetime.fromisoformat(value.replace('Z','+00:00')).astimezone(ZoneInfo('Asia/Ho_Chi_Minh')).strftime('%d/%m/%Y')
    except (ValueError,AttributeError):return str(value)
def filename(courses):
    names=list(dict.fromkeys(c['name'] for c in courses));label=('LOP_' if len(names)==1 else 'CAC_LOP_')+'_'.join(ascii_name(n) for n in names[:5])
    if len(names)>5:label+='_VA_'+str(len(names)-5)+'_LOP_KHAC'
    if len(courses)==1:label+='_'+ascii_name(courses[0]['subject'])
    elif len(names)==1:label+='_CAC_MON'
    return ('BANG_DIEM_'+label)[:190]+'.xlsx'
def build_gradebook(state,course_id,usernames):
    courses=[c for c in state['courses'] if course_id=='all' or c['id']==course_id]
    if not courses:raise ValueError('Không có lớp / môn trong phạm vi tài khoản.')
    book=Workbook();book.remove(book.active);used=set()
    thin=Side(style='thin',color='D1D5DB');border=Border(left=thin,right=thin,top=thin,bottom=thin)
    for course in courses:
        cid=course['id'];base=re.sub(r'[\\/*?:\[\]]',' ',course['name']+' '+course['subject']).upper()[:31] or 'LOP';title=base;i=1
        while title in used:i+=1;title=base[:25]+' ('+str(i)+')'
        used.add(title);sheet=book.create_sheet(title);sheet.sheet_view.showGridLines=False
        ids={e['member_id'] for e in state['enrollments'] if e['course_id']==cid}
        students=sorted([m for m in state['members'] if m['role']=='student' and m['id'] in ids],key=lambda m:ascii_name(m['name']))
        records=[r for r in state['records'] if r['course_id']==cid];columns=[];manual={}
        for r in records:
            d=r['data']
            if r['kind']=='assignment':columns.append({'id':r['id'],'title':d['title'],'date':day(d.get('opensAt') or r['created_at']),'created':r['created_at'],'kind':'exam'})
            elif r['kind']=='grade':
                date=day(d.get('date') or r['created_at']);key=(d['title'].strip().casefold(),date)
                if key not in manual:
                    manual[key]={'id':key,'title':d['title'],'date':date,'created':r['created_at'],'kind':'manual','records':[]};columns.append(manual[key])
                manual[key]['records'].append(r)
        columns.sort(key=lambda c:c['created'])
        headers=['STT','SBD / TÊN ĐĂNG NHẬP','HỌ VÀ TÊN']+[(c['title']+'\n'+c['date']).upper() for c in columns]+['ĐIỂM TRUNG BÌNH'];last=len(headers);end=get_column_letter(last)
        class_name=re.sub(r'^lớp\s*','',course['name'],flags=re.I)
        for row,text in [(1,('BẢNG ĐIỂM - LỚP '+class_name+' - MÔN '+course['subject']).upper()),(2,'ĐIỂM CAO NHẤT MỖI BÀI · THANG ĐIỂM 10'),(3,'SỐ HỌC SINH: '+str(len(students))+' · NGÀY XUẤT: '+datetime.now(ZoneInfo('Asia/Ho_Chi_Minh')).strftime('%d/%m/%Y %H:%M'))]:
            sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=last);cell=sheet.cell(row,1,text);cell.font=Font(name='Times New Roman',size=16 if row==1 else 11,bold=row==1,color='115E59' if row==1 else '64748B',italic=row!=1);cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True);sheet.row_dimensions[row].height=34 if row==1 else 24
        for i,label in enumerate(headers,1):
            cell=sheet.cell(5,i,label);cell.data_type='s';cell.font=Font(name='Times New Roman',size=11,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='0F766E');cell.border=border;cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            sheet.column_dimensions[get_column_letter(i)].width=6 if i==1 else 22 if i==2 else 32 if i==3 else 18
        sheet.row_dimensions[5].height=66
        for n,student in enumerate(students,1):
            scores=[]
            for column in columns:
                if column['kind']=='exam':
                    sub=next((s for s in state['submissions'] if s['assignment_id']==column['id'] and s['member_id']==student['id']),None)
                    result=with_best_result(sub['data'],sub['created_at'])['bestScore'] if sub else None
                    scores.append(result if valid_score(result) else None)
                else:
                    grades=sorted([r for r in column['records'] if r['data']['memberId']==student['id']],key=lambda r:(r['created_at'],r['id']))
                    value=grades[-1]['data']['score'] if grades else None;scores.append(value if valid_score(value) else None)
            actual=[s for s in scores if valid_score(s)];average=round(sum(actual)/len(actual),2) if actual else None
            values=[n,usernames.get(student['id'],student.get('username','')),student['name']]+scores+[average];r=n+5
            for c,value in enumerate(values,1):
                cell=sheet.cell(r,c,value)
                if isinstance(value,str):cell.data_type='s'
                cell.font=Font(name='Times New Roman',size=11);cell.border=border;cell.alignment=Alignment(horizontal='left' if c==3 else 'center',vertical='center',wrap_text=True)
                if n%2==0:cell.fill=PatternFill('solid',fgColor='F0FDFA')
                cell.number_format='@' if c in [2,3] else '0.00' if c>3 else '0'
            sheet.row_dimensions[r].height=28
        sheet.freeze_panes='D6';sheet.auto_filter.ref='A5:'+end+str(max(5,5+len(students)));sheet.print_title_rows='5:5';sheet.print_options.horizontalCentered=True;sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0;sheet.print_area='A1:'+end+str(max(5,5+len(students)))
    out=io.BytesIO();book.save(out);out.seek(0);return out,filename(courses)
