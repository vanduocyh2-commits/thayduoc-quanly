import io
from openpyxl import load_workbook
from gradebook_export import build_gradebook
from test_school import app,client,setup,action,login,state
from test_student_accounts import prepare

def test_excel_class_titles_formats_scores_zero_and_highest():
 state={'courses':[{'id':'c','name':'8A','subject':'Toán'},{'id':'d','name':'8B','subject':'Văn'}],'members':[{'id':'s','name':'=An','role':'student'}],'enrollments':[{'member_id':'s','course_id':'c'}], 'records':[{'id':'a','kind':'assignment','course_id':'c','created_at':'2026-10-01T01:00:00Z','data':{'title':'Bài thi'}},{'id':'g','kind':'grade','course_id':'c','created_at':'2026-10-02T01:00:00Z','data':{'title':'Miệng','score':0,'memberId':'s'}}], 'submissions':[{'assignment_id':'a','member_id':'s','created_at':'2026-10-03','data':{'score':3,'attemptCount':2,'history':[{'score':8,'attemptCount':1}]}}]}
 output,name=build_gradebook(state,'all',{'s':'0001'});assert '8A_8B' in name and name==name.upper().replace('.XLSX','.xlsx')
 book=load_workbook(output);assert len(book.worksheets)==2;s=book.worksheets[0]
 assert s['A1'].value=='BẢNG ĐIỂM - LỚP 8A - MÔN TOÁN';assert s['A1'].alignment.horizontal=='center';assert s['B6'].value=='0001' and s['B6'].data_type=='s';assert s['C6'].data_type=='s';assert [s.cell(6,c).value for c in [4,5,6]]==[8,0,4]
 assert s.freeze_panes=='D6' and s.page_setup.orientation=='landscape';assert s['D5'].value=='BÀI THI\n01/10/2026'

def test_gradebook_endpoint_is_class_scoped_and_auth_required(client,app):
 t1,t2,s1,s2,ids,cs=prepare(client,app)
 assert action(t1,{'action':'record','kind':'grade','courseId':cs['10A'],'data':{'title':'Bài 1','score':9,'memberId':ids['student1']}}).status_code==200
 result=t1.get('/api/grades/export?courseId='+cs['10A']);assert result.status_code==200
 assert 'BANG_DIEM_LOP_10A_TOAN.xlsx' in result.headers['Content-Disposition']
 assert load_workbook(io.BytesIO(result.data)).active['D6'].value==9
 assert t1.get('/api/grades/export?courseId='+cs['10B']).status_code==403
 assert app.test_client().get('/api/grades/export').status_code==401
 personal=s1.get('/api/grades/export');assert personal.status_code==200
 assert load_workbook(io.BytesIO(personal.data)).active.max_row==6
