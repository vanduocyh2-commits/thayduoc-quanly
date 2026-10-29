"""Read TN Maker's question-by-code matrix and the older row-based template."""
import io
import re
import unicodedata
import zipfile
from openpyxl import load_workbook


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def header_text(value):
    value = unicodedata.normalize('NFD', cell_text(value).lower()).replace('đ', 'd')
    return ''.join(c for c in value if c.isalnum() and not unicodedata.combining(c))


def read_answer_keys(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if sum(info.file_size for info in archive.infolist()) > 20 * 1024 * 1024:
            raise ValueError('File đáp án sau giải nén quá lớn.')
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        values = list(workbook.active.iter_rows(max_row=1002, max_col=12, values_only=True))
    finally:
        workbook.close()
    while values and not any(v is not None for v in values[-1]):
        values.pop()
    if not values:
        raise ValueError('File đáp án chưa có dữ liệu.')
    if len(values) > 1001:
        raise ValueError('Tối đa 1000 dòng đáp án.')
    if [header_text(v) for v in values[0][:4]] == ['made', 'dang', 'cau', 'dapan']:
        rows = []
        for row in values[1:]:
            if not any(v is not None for v in row):
                continue
            code, kind, number, answer = row[:4]
            if kind not in ['quiz', 'truefalse', 'short'] or type(number) not in [int, float] or int(number) != number:
                raise ValueError('Dạng hoặc số câu chưa đúng mẫu.')
            rows.append(dict(code=cell_text(code), type=kind, number=int(number), answer=cell_text(answer)))
        if not rows:
            raise ValueError('File đáp án chưa có dữ liệu.')
        return rows
    if header_text(values[0][0]) != 'caumade':
        raise ValueError('Dùng file TN Maker: ô đầu là Câu\\Mã Đề, các cột tiếp theo là mã đề.')
    last = max((i for i, v in enumerate(values[0]) if cell_text(v)), default=0)
    codes = [cell_text(v) for v in values[0][1:last + 1]]
    if not 1 <= len(codes) <= 10 or any(not re.fullmatch(r'\d{3,4}', code) for code in codes):
        raise ValueError('TN Maker cần 1–10 cột mã đề, mỗi mã gồm 3–4 chữ số.')
    if len(set(codes)) != len(codes):
        raise ValueError('Mã đề trong file đáp án bị trùng.')
    result = []
    section = None
    previous = 0
    seen = set()
    order = {'quiz': 0, 'truefalse': 1, 'short': 2}
    for line, row in enumerate(values[1:], start=2):
        if not any(v is not None for v in row):
            continue
        number = cell_text(row[0])
        if not re.fullmatch(r'\d{1,3}', number) or not 1 <= int(number) <= 100:
            raise ValueError(f'Dòng {line}: số câu phải là số nguyên từ 1 đến 100.')
        number = int(number)
        answers = [cell_text(v) for v in row[1:len(codes) + 1]]
        if any(not answer for answer in answers):
            code = codes[next(i for i, answer in enumerate(answers) if not answer)]
            raise ValueError(f'Dòng {line}, mã {code}: còn thiếu đáp án.')
        if any(cell_text(v) for v in row[len(codes) + 1:]):
            raise ValueError(f'Dòng {line}: có đáp án ngoài các cột mã đề.')
        if section is None or number == 1:
            compact = [re.sub(r'[\s,;]+', '', answer.upper()).replace('D', 'Đ') for answer in answers]
            kind = 'quiz' if all(re.fullmatch('[ABCD]', a.upper()) for a in answers) else 'truefalse' if all(re.fullmatch('[ĐS]{4}', a) for a in compact) else 'short'
            if kind in seen or section is not None and order[kind] <= order[section]:
                raise ValueError(f'Dòng {line}: các phần TN Maker cần theo thứ tự trắc nghiệm, đúng/sai, trả lời ngắn; mỗi phần bắt đầu từ câu 1.')
            section, previous = kind, 0
            seen.add(kind)
        if number != previous + 1:
            raise ValueError(f'Dòng {line}: số câu trong mỗi phần phải liên tục, bắt đầu từ 1.')
        previous = number
        for code, answer in zip(codes, answers):
            if section == 'quiz':
                answer = answer.upper()
                if not re.fullmatch('[ABCD]', answer):
                    raise ValueError(f'Dòng {line}, mã {code}: trắc nghiệm cần A, B, C hoặc D.')
            elif section == 'truefalse':
                answer = re.sub(r'[\s,;]+', '', answer.upper()).replace('D', 'Đ')
                if not re.fullmatch('[ĐS]{4}', answer):
                    raise ValueError(f'Dòng {line}, mã {code}: đúng/sai cần đủ 4 ý Đ/S.')
                answer = ' '.join(answer)
            if len(answer) > 200:
                raise ValueError(f'Dòng {line}, mã {code}: đáp án tối đa 200 ký tự.')
            result.append(dict(code=code, type=section, number=number, answer=answer))
    if not result:
        raise ValueError('File đáp án chưa có dữ liệu.')
    if len(result) > 1000 or len(result) // len(codes) > 100:
        raise ValueError('Tối đa 100 câu mỗi mã đề và 1000 đáp án.')
    return result
