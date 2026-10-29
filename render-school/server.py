"""Independent school portal for Render. No ChatGPT or Cloudflare dependency."""
import os,json,re,hmac,secrets,hashlib,time,io,unicodedata,zipfile,math
from collections import Counter
from pathlib import Path
from datetime import datetime,timezone
from functools import wraps
from flask import Flask,request,jsonify,g,send_from_directory,send_file,make_response
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash,check_password_hash
from werkzeug.middleware.proxy_fix import ProxyFix
from sqlalchemy import create_engine,MetaData,Table,Column,String,Integer,Text,LargeBinary,ForeignKey,UniqueConstraint,select,insert,update,delete,and_,func
from sqlalchemy.exc import IntegrityError
from sqlalchemy import text as sql_text
from answer_keys import read_answer_keys
from pdf_conversion import convert_docx_to_pdf,ConversionError
from exam_results import with_best_result,grade_attempt
from school_uploads import read_attachment

metadata=MetaData()
members=Table('school_members',metadata,Column('id',String,primary_key=True),Column('username',String,unique=True,nullable=False),Column('password_hash',Text,nullable=False),Column('name',String,nullable=False),Column('email',String),Column('role',String,nullable=False),Column('disabled',Integer,nullable=False,default=0))
courses=Table('school_courses',metadata,Column('id',String,primary_key=True),Column('name',String,nullable=False),Column('subject',String,nullable=False),Column('teacher_id',String,ForeignKey('school_members.id'),nullable=False))
enrollments=Table('school_enrollments',metadata,Column('course_id',String,ForeignKey('school_courses.id'),primary_key=True),Column('member_id',String,ForeignKey('school_members.id'),primary_key=True))
records=Table('school_records',metadata,Column('id',String,primary_key=True),Column('course_id',String,ForeignKey('school_courses.id'),nullable=False),Column('kind',String,nullable=False),Column('author_id',String,ForeignKey('school_members.id'),nullable=False),Column('data',Text,nullable=False),Column('created_at',String,nullable=False))
submissions=Table('school_submissions',metadata,Column('id',String,primary_key=True),Column('assignment_id',String,ForeignKey('school_records.id'),nullable=False),Column('member_id',String,ForeignKey('school_members.id'),nullable=False),Column('data',Text,nullable=False),Column('created_at',String,nullable=False),UniqueConstraint('assignment_id','member_id'))
sessions=Table('school_sessions',metadata,Column('token_hash',String,primary_key=True),Column('member_id',String,ForeignKey('school_members.id'),nullable=False),Column('csrf',String,nullable=False),Column('expires_at',Integer,nullable=False))
limits=Table('school_login_limits',metadata,Column('key',String,primary_key=True),Column('failures',Integer,nullable=False),Column('expires_at',Integer,nullable=False))

library_files=Table('school_library_files',metadata,Column('id',String,primary_key=True),Column('record_id',String,ForeignKey('school_records.id'),unique=True,nullable=False),Column('name',String,nullable=False),Column('mime',String,nullable=False),Column('size',Integer,nullable=False),Column('content',LargeBinary,nullable=False))

exam_files=Table('school_exam_variant_files',metadata,Column('id',String,primary_key=True),Column('record_id',String,ForeignKey('school_records.id'),nullable=False),Column('code',String,nullable=False),Column('name',String,nullable=False),Column('mime',String,nullable=False),Column('size',Integer,nullable=False),Column('content',LargeBinary,nullable=False),UniqueConstraint('record_id','code'))

rosters=Table('school_rosters',metadata,Column('id',String,primary_key=True),Column('course_id',String,ForeignKey('school_courses.id'),nullable=False),Column('author_id',String,ForeignKey('school_members.id'),nullable=False),Column('data',Text,nullable=False),Column('status',String,nullable=False),Column('created_at',String,nullable=False))
submission_files=Table('school_submission_files',metadata,Column('id',String,primary_key=True),Column('submission_id',String,ForeignKey('school_submissions.id'),nullable=False),Column('attempt_count',Integer,nullable=False),Column('name',String,nullable=False),Column('mime',String,nullable=False),Column('size',Integer,nullable=False),Column('content',LargeBinary,nullable=False))
exam_resources=Table('school_exam_resources',metadata,Column('id',String,primary_key=True),Column('record_id',String,ForeignKey('school_records.id'),nullable=False),Column('kind',String,nullable=False),Column('name',String,nullable=False),Column('mime',String,nullable=False),Column('size',Integer,nullable=False),Column('content',LargeBinary,nullable=False),UniqueConstraint('record_id','kind'))

class AppError(Exception):
 def __init__(self,message,status=400):self.message=message;self.status=status

def fail(message,status=400):raise AppError(message,status)
def text(v,maxlen=200,required=True):
 if not isinstance(v,str) or len(v)>maxlen or (required and not v.strip()):fail('Nội dung chưa hợp lệ.')
 return v.strip()
def normalized_short(value,ignore_case=False):
 value=' '.join(value.strip().split())
 if re.fullmatch(r'[+-]?(?:[0-9]+(?:[.,][0-9]+)?|[.,][0-9]+)(?:[eE][+-]?[0-9]+)?',value):value=value.replace(',', '.')
 return value.casefold() if ignore_case else value
def validate_auto_question(q):
 if not isinstance(q,dict):fail('Câu hỏi chưa hợp lệ.')
 text(q.get('text'),2000)
 kind=q.get('type','quiz')
 if kind not in ['quiz','truefalse','short']:fail('Dạng câu hỏi chưa hợp lệ.')
 if kind in ['quiz','truefalse']:
  if not isinstance(q.get('answers'),list) or len(q['answers'])!=4:fail('Cần đủ 4 phương án / ý.')
  for answer in q['answers']:text(answer,1000)
 if kind=='quiz' and (type(q.get('correct')) is not int or not 0<=q['correct']<=3):fail('Chọn đáp án A–D.')
 if kind=='truefalse' and (not isinstance(q.get('correct'),list) or len(q['correct'])!=4 or any(type(v) is not bool for v in q['correct'])):fail('Cần đáp án Đúng / Sai cho đủ 4 ý.')
 if kind=='short':
  text(q.get('correct'),200)
  if type(q.get('ignoreCase',False)) is not bool:fail('Cấu hình trả lời ngắn không hợp lệ.')
def auto_question_score(q,answer,mode,weighted):
 kind=q.get('type','quiz');points=q.get('points',0) if weighted else 1
 if kind=='truefalse':
  if not isinstance(answer,list) or len(answer)!=4 or any(type(v) is not bool for v in answer):fail('Cần trả lời đủ 4 ý Đúng / Sai.')
  right=sum(a==b for a,b in zip(answer,q['correct']))
  return points*(right/4 if mode=='equal' else [0,.1,.25,.5,1][right])
 if kind=='short':
  value=text(answer,200)
  return points if normalized_short(value,q.get('ignoreCase',False))==normalized_short(q['correct'],q.get('ignoreCase',False)) else 0
 if type(answer) is not int or not 0<=answer<=3:fail('Cần chọn đáp án A–D.')
 return points if answer==q['correct'] else 0

def assigned_exam(data,rid,mid):
 variants=data.get('variants',[])
 if not variants:return data
 index=int(hashlib.sha256((rid+':'+mid).encode()).hexdigest(),16)%len(variants)
 v=next((v for v in variants if data.get('codePolicy')=='fixed' and v['code']==data.get('fixedExamCode')),variants[index])
 return {**data,'questions':v['questions'],'essays':v['essays'],'examCode':v['code'],'file':v.get('file',data.get('file'))}
def question_signature(q):
 kind=q.get('type','quiz')
 d={k:q.get(k) for k in ['text','points','topic','level','outcome']};d['type']=kind
 if kind=='quiz':d.update(answers=sorted(q['answers']),correct=q['answers'][q['correct']])
 elif kind=='truefalse':d.update(answers=q['answers'],correct=q['correct'])
 elif kind=='short':d.update(correct=q['correct'],ignoreCase=q.get('ignoreCase',False))
 return json.dumps(d,sort_keys=True,ensure_ascii=False)
def validate_variants(data):
 variants=data.get('variants',[])
 if not isinstance(variants,list) or len(variants)>10:fail('Tối đa 10 mã đề.')
 base=Counter(question_signature(q) for q in data['questions']+data.get('essays',[]));codes=set()
 for v in variants:
  if not isinstance(v,dict) or not isinstance(v.get('code'),str) or not re.fullmatch(r'\d{3,4}',v['code']) or v['code'] in codes:fail('Mã đề phải khác nhau, gồm 3–4 chữ số.')
  codes.add(v['code'])
  if not isinstance(v.get('questions'),list) or not isinstance(v.get('essays'),list) or len(v['questions'])!=len(data['questions']) or len(v['essays'])!=len(data.get('essays',[])):fail('Mã đề phải giữ đủ câu của đề gốc.')
  for q in v['questions']:validate_auto_question(q)
  for q in v['essays']:
   if not isinstance(q,dict) or q.get('type')!='essay':fail('Câu tự luận trong mã đề chưa hợp lệ.')
  if data.get('source')=='uploaded':
   validate_exam_data({**data,'questions':v['questions'],'essays':v['essays'],'variants':[]})
   continue
  try:valid=Counter(question_signature(q) for q in v['questions']+v['essays'])==base
  except (TypeError,KeyError,IndexError):valid=False
  if not valid:fail('Bản đảo phải giữ nguyên nội dung, thang điểm và đáp án đúng của đề gốc.')

def validate_exam_data(d):
 if d.get('type') not in ['essay','quiz','mixed']:fail('Hình thức chưa hợp lệ.')
 if d.get('due'):parse_date(d['due'])
 if d['type'] in ['quiz','mixed']:
  qs=d.get('questions')
  if not isinstance(qs,list) or not 1<=len(qs)<=100:fail('Cần từ 1 đến 100 câu hỏi.')
  for q in qs:validate_auto_question(q)
 else:d['questions']=[]
 essays=d.get('essays',[])
 if not isinstance(essays,list) or len(essays)>100:fail('Tối đa 100 câu tự luận.')
 for q in essays:
  if not isinstance(q,dict):fail('Câu tự luận chưa hợp lệ.')
  text(q.get('text'),2000)
  q['type']='essay'
 if d['type']=='mixed' and not essays:fail('Đề kết hợp cần câu tự luận.')
 if d.get('composed'):
  all_q=d['questions']+essays
  if not 1<=len(all_q)<=100:fail('Đề cần từ 1 đến 100 câu hỏi.')
  for q in all_q:
   if type(q.get('points')) not in [int,float] or not math.isfinite(q['points']) or not 0<q['points']<=10:fail('Điểm câu hỏi phải lớn hơn 0 và không quá 10.')
   q['topic']=text(q.get('topic'),200);q['level']=text(q.get('level'),100)
   q['outcome']=text(q.get('outcome',''),2000,False)
   if q['level'] not in ['Nhận biết','Thông hiểu','Vận dụng','Vận dụng cao']:fail('Mức độ không hợp lệ.')
  if abs(sum(q['points'] for q in all_q)-10)>0.001:fail('Tổng điểm đề phải bằng 10.')
 if d.get('opensAt'):parse_date(d['opensAt'])
 if d.get('opensAt') and d.get('due') and parse_date(d['opensAt'])>=parse_date(d['due']):fail('Giờ kết thúc phải sau giờ bắt đầu.')
 if type(d.get('maxAttempts',1)) is not int or not 0<=d.get('maxAttempts',1)<=100:fail('Số lần làm bài từ 1–100 hoặc 0 để không giới hạn.')
 if d.get('minutes') is not None and (type(d['minutes']) is not int or not 1<=d['minutes']<=300):fail('Thời gian cần từ 1 đến 300 phút.')
 if d.get('grade') is not None and str(d['grade']) not in [str(i) for i in range(1,13)]:fail('Khối lớp không hợp lệ.')
 if d.get('tfScoring','progressive') not in ['progressive','equal']:fail('Cách chấm Đúng / Sai không hợp lệ.')
 d['includeMatrix']=bool(d.get('includeMatrix',False))
 validate_variants(d)
 if d.get('subjectCode') is not None and not re.fullmatch(r'\d{2}',str(d['subjectCode'])):fail('Mã môn cần đúng 2 chữ số.')
 d.pop('file',None)

def exam_upload(file):
 if not file or not file.filename:fail('Hãy chọn đề từ máy.')
 raw=file.stream.read(5*1024*1024+1)
 if not raw or len(raw)>5*1024*1024:fail('Tệp phải có nội dung, tối đa 5 MB.')
 ext=Path(file.filename).suffix.lower();preview=''
 allowed={'.pdf':'application/pdf','.docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','.txt':'text/plain','.md':'text/markdown'}
 if ext not in allowed:fail('Hỗ trợ PDF, Word .docx, TXT và Markdown. File .doc cần lưu lại thành .docx; JSON nhập thành câu hỏi theo mẫu.')
 if ext=='.pdf' and not raw.startswith(b'%PDF-'):fail('Nội dung không phải tệp PDF hợp lệ.')
 if ext=='.docx':
  try:
   from xml.etree import ElementTree as ET
   with zipfile.ZipFile(io.BytesIO(raw)) as z:
    if len(z.infolist())>2000 or sum(i.file_size for i in z.infolist())>20*1024*1024:fail('File Word sau giải nén quá lớn.')
    if '[Content_Types].xml' not in z.namelist() or 'word/document.xml' not in z.namelist() or any('vbaProject' in n for n in z.namelist()):fail('Cần file Word .docx không chứa macro.')
    for relation_file in (n for n in z.namelist() if n.endswith('.rels')):
     relations=ET.fromstring(z.read(relation_file))
     if any(r.get('TargetMode')=='External' and not r.get('Type','').endswith('/hyperlink') for r in relations):fail('Word có hình ảnh hoặc dữ liệu liên kết ngoài. Hãy nhúng nội dung vào file trước khi tải.')
    root=ET.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    preview='\n'.join(''.join(n.text or '' for n in paragraph.iter() if n.tag.endswith('}t')) for paragraph in root.findall('.//w:p',ns))[:20000]
  except AppError:raise
  except Exception:fail('Không đọc được Word. Hãy kiểm tra tệp .docx không có mật khẩu.')
 if ext in ['.txt','.md']:
  try:preview=raw.decode('utf-8-sig')
  except UnicodeDecodeError:fail('File văn bản cần mã hóa UTF-8.')
  if '\x00' in preview:fail('File không phải văn bản hợp lệ.')
  preview=preview[:20000]
 name=secure_filename(file.filename) or 'de-thi'+ext
 if ext=='.docx':
  try:raw=convert_docx_to_pdf(raw)
  except ConversionError as error:fail(str(error),503)
  name=str(Path(name).with_suffix('.pdf'));ext='.pdf';preview=''
 return {'name':name,'mime':allowed[ext],'content':raw,'size':len(raw),'preview':preview}
def outline_data(p):
 if not isinstance(p,dict):fail('Đề cương chưa hợp lệ.')
 title=text(p.get('title'),200);grade=str(p.get('grade',''))
 if grade not in [str(i) for i in range(1,13)]:fail('Chọn khối lớp từ 1 đến 12.')
 sections=p.get('sections')
 if not isinstance(sections,list) or not 1<=len(sections)<=50:fail('Đề cương cần từ 1 đến 50 chủ đề.')
 out=[]
 for section in sections:
  if not isinstance(section,dict):fail('Chủ đề chưa hợp lệ.')
  value={'topic':text(section.get('topic'),200)}
  for k in ['heading','chapter','lesson']:value[k]=text(section.get(k,''),200,False)
  for k in ['objectives','theory','exercises']:value[k]=text(section.get(k,''),10000,False)
  if not any(value[k] for k in ['objectives','theory','exercises']):fail('Mỗi chủ đề cần nội dung ôn tập hoặc bài tập.')
  out.append(value)
 if sum(len(v) for a in out for v in a.values())>100000:fail('Đề cương quá dài; hãy tách thành nhiều phần.')
 matrix=p.get('sourceMatrix',[])
 if not isinstance(matrix,list) or len(matrix)>100:fail('Ma trận nguồn không hợp lệ.')
 cleaned=[]
 for r in matrix:
  if not isinstance(r,dict) or r.get('type') not in ['quiz','truefalse','short','essay']:fail('Dạng câu trong ma trận không hợp lệ.')
  counts=r.get('counts')
  if not isinstance(counts,list) or len(counts)!=4 or any(type(n) is not int or not 0<=n<=100 for n in counts):fail('Phân bố ma trận không hợp lệ.')
  points=r.get('points',0)
  if type(points) not in [int,float] or not math.isfinite(points) or not 0<=points<=10:fail('Điểm ma trận không hợp lệ.')
  cleaned.append(dict(type=r['type'],topic=text(r.get('topic'),200),counts=counts,count=sum(counts),points=points))
 return dict(title=title,grade=grade,body=text(p.get('body',''),10000,False),sections=out,published=p.get('published') is True,inLibrary=p.get('inLibrary',True) is True,requirements=text(p.get('requirements',''),10000,False),sourceMatrix=cleaned)

def password(v):
 if not isinstance(v,str) or not 10<=len(v)<=200:fail('Mật khẩu cần từ 10 đến 200 ký tự.')
 return generate_password_hash(v)
def username(v):
 v=text(v,60).lower()
 if not re.fullmatch(r'[a-z0-9_.-]{3,60}',v):fail('Tên đăng nhập cần 3–60 chữ không dấu, số, dấu chấm, gạch dưới hoặc gạch nối.')
 return v
def uid():return secrets.token_hex(16)
def now():return datetime.now(timezone.utc).isoformat()
def parse_date(value):
 try:return datetime.fromisoformat(value.replace('Z','+00:00')).timestamp()
 except (ValueError,AttributeError,TypeError):fail('Ngày giờ chưa hợp lệ.')
def may_teach(m,c):return not m['disabled'] and (m['role']=='admin' or (m['role']=='teacher' and c['teacher_id']==m['id']))
def public_member(m):return {k:v for k,v in m.items() if k!='password_hash'}

def normalized(v):return ''.join(c for c in unicodedata.normalize('NFD',str(v).lower().replace('đ','d')) if not unicodedata.combining(c)).strip()
def roster_students(values):
 if not isinstance(values,list) or not 1<=len(values)<=500:fail('Danh sách cần từ 1 đến 500 học sinh.')
 result=[];seen=set()
 for i,v in enumerate(values,1):
  if not isinstance(v,dict):fail('Dòng học sinh chưa hợp lệ.')
  code=username(v.get('code'));name=text(v.get('name'),200)
  if code in seen:fail(f'Dòng {i}: trùng mã học sinh {code}.')
  seen.add(code);result.append({'code':code,'name':name})
 return result

def parse_roster(upload):
 from openpyxl import load_workbook
 if not upload:fail('Hãy chọn file Excel .xlsx.')
 if not upload.filename.lower().endswith('.xlsx'):fail('Hỗ trợ Excel .xlsx. Với file .xls, hãy lưu thành .xlsx trước.')
 raw=upload.read()
 try:
  with zipfile.ZipFile(io.BytesIO(raw)) as z:
   if sum(x.file_size for x in z.infolist())>40*1024*1024:fail('File Excel sau giải nén quá lớn.')
  wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
 except AppError:raise
 except Exception:fail('Không đọc được file Excel. Hãy kiểm tra file .xlsx không có mật khẩu.')
 result=[];errors=[];seen=set();found=False
 try:
  for sheet in wb:
   header=None
   for line,row in enumerate(sheet.iter_rows(max_row=min(sheet.max_row or 0,2000),max_col=50,values_only=True),1):
    cells=[normalized(v or '') for v in row]
    if header is None:
     code=next((i for i,v in enumerate(cells) if v in ['ma hoc sinh','ma hs','ma dinh danh bo gd&dt','ma dinh danh','sbd','ten dang nhap']),None)
     name=next((i for i,v in enumerate(cells) if v in ['ho va ten','ho ten','ho va ten hoc sinh']),None)
     if code is not None and name is not None:header=(code,name);found=True
     continue
    code_value=row[header[0]];name_value=row[header[1]]
    if not code_value and not name_value:continue
    # Skip the second header line and footer summaries in Ministry grade exports.
    if not code_value and (not row[0] or not str(row[0]).strip().isdigit()):continue
    code=str(int(code_value)) if isinstance(code_value,(int,float)) and float(code_value).is_integer() else str(code_value or '').strip()
    try:
     item=roster_students([{'code':code,'name':str(name_value or '')}])[0]
     if item['code'] in seen:fail('Trùng mã học sinh trong file.')
     seen.add(item['code']);result.append(item)
    except AppError as e:errors.append(f'{sheet.title}, dòng {line}: {e.message}')
    if len(result)>500:fail('Mỗi lần tải tối đa 500 học sinh.')
   if header and (sheet.max_row or 0)>2000:errors.append(f'{sheet.title}: bảng vượt quá 2000 dòng, hãy tách file.')
 finally:wb.close()
 if not found:fail('Không tìm thấy cột Mã định danh Bộ GD&ĐT / Mã học sinh và Họ và tên trong 50 cột đầu.')
 if not result and not errors:fail('Không có học sinh trong file.')
 return result,errors

def create_app(config=None):
 app=Flask(__name__,static_folder='static',static_url_path='/static');app.config.update(MAX_CONTENT_LENGTH=26*1024*1024,SECRET_KEY=os.getenv('SECRET_KEY'),DATABASE_URL=os.getenv('DATABASE_URL'),ADMIN_SETUP_TOKEN=os.getenv('ADMIN_SETUP_TOKEN'),SESSION_SECONDS=43200,TESTING=False)
 if config:app.config.update(config)
 if not app.config['SECRET_KEY'] or len(app.config['SECRET_KEY'])<32:raise RuntimeError('Configure SECRET_KEY (at least 32 characters).')
 url=app.config['DATABASE_URL']
 if not url:raise RuntimeError('Configure DATABASE_URL for a separate persistent school database.')
 if url.startswith('postgres://'):url='postgresql+psycopg://'+url[len('postgres://'):]
 elif url.startswith('postgresql://'):url='postgresql+psycopg://'+url[len('postgresql://'):]
 engine=create_engine(url,pool_pre_ping=True,**({'connect_args':{'check_same_thread':False}} if url.startswith('sqlite') else {'pool_size':5,'max_overflow':5}))
 app.extensions['school_engine']=engine
 if engine.dialect.name=='postgresql':
  with engine.begin() as conn:
   conn.execute(sql_text('SELECT pg_advisory_xact_lock(719246830)'))
   metadata.create_all(conn)
   for table in metadata.sorted_tables:
    conn.execute(sql_text(f'ALTER TABLE {table.name} ENABLE ROW LEVEL SECURITY'))
    conn.execute(sql_text(f'REVOKE ALL ON TABLE {table.name} FROM PUBLIC'))
 else:metadata.create_all(engine)
 app.wsgi_app=ProxyFix(app.wsgi_app,x_for=1,x_proto=1,x_host=0)
 cookie='school_session'
 def one(table,where):return g.db.execute(select(table).where(where)).mappings().first()
 def rows(table):return [dict(x) for x in g.db.execute(select(table)).mappings()]
 def stored_file_bytes():return sum(g.db.execute(select(func.coalesce(func.sum(t.c.size),0))).scalar_one() for t in [library_files,exam_files,submission_files,exam_resources])
 def origin():
  if request.headers.get('Origin')!=request.host_url.rstrip('/'):fail('Nguồn yêu cầu không hợp lệ.',403)
 def body():
  if not request.is_json:fail('Yêu cầu JSON.',415)
  p=request.get_json(silent=True)
  if not isinstance(p,dict):fail('Nội dung chưa hợp lệ.')
  return p
 def admin():
  if g.member['role']!='admin':fail('Chỉ quản trị viên được thao tác.',403)
 def get_course(cid):
  c=one(courses,courses.c.id==text(cid))
  if not c:fail('Không tìm thấy lớp / môn.',404)
  return c
 def teach(cid):
  c=get_course(cid)
  if not may_teach(g.member,c):fail('Bạn không được phân công lớp / môn này.',403)
  return c
 def read(cid):
  c=get_course(cid)
  if not may_teach(g.member,c) and not (g.member['role']=='student' and one(enrollments,and_(enrollments.c.course_id==cid,enrollments.c.member_id==g.member['id']))):fail('Bạn không thuộc lớp / môn này.',403)
  return c
 def require_member():
  if not g.member:fail('Bạn cần đăng nhập.',401)
 def reactivate_removed_student(cid,mid):
  rid='removed:'+cid+':'+mid;r=one(records,records.c.id==rid)
  if r:
   d=json.loads(r['data']);d['restoredAt']=now();g.db.execute(update(records).where(records.c.id==rid).values(data=json.dumps(d,ensure_ascii=False)))
 def remove_roster_student(cid,mid):
  teach(cid);g.db.execute(select(courses).where(courses.c.id==cid).with_for_update()).first()
  pupil=one(members,and_(members.c.id==mid,members.c.role=='student'))
  where=and_(enrollments.c.course_id==cid,enrollments.c.member_id==mid)
  if not pupil or not one(enrollments,where):fail('Học sinh không còn trong lớp / môn này.',404)
  backups=[]
  for roster in g.db.execute(select(rosters).where(rosters.c.course_id==cid).with_for_update()).mappings():
   students=json.loads(roster['data']);matches=[st for st in students if st.get('code')==pupil['username']]
   if not matches:continue
   remaining=[st for st in students if st.get('code')!=pupil['username']]
   backups.append(dict(id=roster['id'],students=matches,status=roster['status']))
   g.db.execute(update(rosters).where(rosters.c.id==roster['id']).values(data=json.dumps(remaining,ensure_ascii=False),status='rejected' if not remaining and roster['status']=='pending' else roster['status']))
  data=dict(memberId=mid,name=pupil['name'],code=pupil['username'],removedAt=now(),removedBy=g.member['id'],rosterRows=backups)
  rid='removed:'+cid+':'+mid;values=dict(course_id=cid,kind='rosterRemoval',author_id=g.member['id'],created_at=now(),data=json.dumps(data,ensure_ascii=False))
  if one(records,records.c.id==rid):g.db.execute(update(records).where(records.c.id==rid).values(**values))
  else:g.db.execute(insert(records).values(id=rid,**values))
  g.db.execute(delete(enrollments).where(where))
 def restore_roster_student(rid):
  r=one(records,and_(records.c.id==rid,records.c.kind=='rosterRemoval'))
  if not r:fail('Không tìm thấy học sinh đã xóa.',404)
  teach(r['course_id']);g.db.execute(select(courses).where(courses.c.id==r['course_id']).with_for_update()).first()
  r=g.db.execute(select(records).where(records.c.id==rid).with_for_update()).mappings().first();d=json.loads(r['data'])
  if d.get('restoredAt'):fail('Học sinh đã được khôi phục.',409)
  pupil=one(members,and_(members.c.id==d['memberId'],members.c.role=='student',members.c.disabled==0))
  if not pupil:fail('Tài khoản học sinh đang khóa hoặc không tồn tại. Nhờ quản trị viên kiểm tra.',409)
  where=and_(enrollments.c.course_id==r['course_id'],enrollments.c.member_id==pupil['id'])
  if not one(enrollments,where):g.db.execute(insert(enrollments).values(course_id=r['course_id'],member_id=pupil['id']))
  for backup in d.get('rosterRows',[]):
   roster=one(rosters,and_(rosters.c.id==backup['id'],rosters.c.course_id==r['course_id']))
   if not roster:continue
   students=json.loads(roster['data'])
   for st in backup['students']:
    if not any(row.get('code')==st['code'] for row in students):students.append(st)
   status=backup['status'] if not json.loads(roster['data']) and roster['status']=='rejected' else roster['status']
   g.db.execute(update(rosters).where(rosters.c.id==roster['id']).values(data=json.dumps(students,ensure_ascii=False),status=status))
  reactivate_removed_student(r['course_id'],pupil['id'])
 def ratekeys(user):
  return [('ip:'+request.remote_addr,30),('user:'+user,10)]
 def check_rate(user):
  for key,threshold in ratekeys(user):
   r=one(limits,limits.c.key==hashlib.sha256(key.encode()).hexdigest())
   if r and r['expires_at']>int(time.time()) and r['failures']>=threshold:fail('Thử đăng nhập quá nhiều lần. Vui lòng chờ 15 phút.',429)
 def record_failure(user):
  # Atomic increment for both PostgreSQL and SQLite, including concurrent workers.
  for key,_ in ratekeys(user):
   key=hashlib.sha256(key.encode()).hexdigest();ts=int(time.time());r=one(limits,limits.c.key==key)
   if r and r['expires_at']>ts:g.db.execute(update(limits).where(limits.c.key==key).values(failures=limits.c.failures+1))
   elif r:g.db.execute(update(limits).where(limits.c.key==key).values(failures=1,expires_at=ts+900))
   else:
    dialect=engine.dialect.name
    from sqlalchemy.dialects.postgresql import insert as pg_insert
    from sqlalchemy.dialects.sqlite import insert as sq_insert
    ins=(pg_insert if dialect=='postgresql' else sq_insert)(limits).values(key=key,failures=1,expires_at=ts+900)
    g.db.execute(ins.on_conflict_do_update(index_elements=['key'],set_={'failures':limits.c.failures+1}))
  g.db.commit()
 def sign_in(m):
  token=secrets.token_urlsafe(32);csrf=secrets.token_urlsafe(32);ts=int(time.time())
  g.db.execute(delete(sessions).where(sessions.c.expires_at<ts));g.db.execute(insert(sessions).values(token_hash=hashlib.sha256(token.encode()).hexdigest(),member_id=m['id'],csrf=csrf,expires_at=ts+app.config['SESSION_SECONDS']))
  res=jsonify(ok=True);res.set_cookie(cookie,token,httponly=True,secure=not app.config['TESTING'],samesite='Lax',max_age=app.config['SESSION_SECONDS'],path='/');return res
 @app.before_request
 def before():
  if not request.path.startswith('/api/') and request.path!='/healthz':return
  g.db=engine.connect();g.member=None;g.auth=None
  if request.path in ['/api/school','/api/login','/api/setup'] and (request.content_length or 0)>160000:fail('Nội dung quá lớn.',413)
  token=request.cookies.get(cookie)
  if token:
   sess=one(sessions,and_(sessions.c.token_hash==hashlib.sha256(token.encode()).hexdigest(),sessions.c.expires_at>int(time.time())))
   if sess:
    m=one(members,and_(members.c.id==sess['member_id'],members.c.disabled==0))
    if m:g.member=m;g.auth=sess
  if request.method=='POST':
   origin()
   if request.path not in ['/api/login','/api/setup']:
    require_member()
    if not hmac.compare_digest(request.headers.get('X-CSRF-Token',''),g.auth['csrf']):fail('Phiên xác thực không hợp lệ. Hãy tải lại trang.',403)
 @app.after_request
 def after(res):
  if hasattr(g,'db'):
   if res.status_code<400:g.db.commit()
   else:g.db.rollback()
  res.headers['X-Content-Type-Options']='nosniff';res.headers['Referrer-Policy']='same-origin';res.headers['X-Frame-Options']='DENY'
  res.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; frame-src 'self' blob:; base-uri 'self'; form-action 'self'"
  if request.path.startswith(('/api/exams/','/api/submissions/')) and request.args.get('preview')=='1' and res.mimetype=='application/pdf':
   res.headers['X-Frame-Options']='SAMEORIGIN';res.headers['Content-Security-Policy']="default-src 'none'; frame-ancestors 'self'"
  if request.path.startswith('/api/'):res.headers['Cache-Control']='no-store'
  return res
 @app.teardown_request
 def close(_):
  if hasattr(g,'db'):g.db.close()
 @app.errorhandler(AppError)
 def app_error(e):return jsonify(error=e.message),e.status
 @app.errorhandler(IntegrityError)
 def conflict(_):return jsonify(error='Tên đăng nhập, ghi danh hoặc bài nộp đã tồn tại.'),409
 @app.errorhandler(413)
 def too_large(_):return jsonify(error='Nội dung quá lớn.'),413
 @app.errorhandler(500)
 def generic(_):return jsonify(error='Máy chủ gặp lỗi. Vui lòng thử lại.'),500
 @app.get('/healthz')
 def health():g.db.execute(select(1));return jsonify(status='ok')
 @app.get('/api/auth')
 def auth():return jsonify(csrf=g.auth['csrf'] if g.auth else '',setupNeeded=not bool(one(members,members.c.role=='admin')))
 @app.post('/api/setup')
 def setup():
  if one(members,members.c.role=='admin'):fail('Quản trị viên đã được thiết lập.',409)
  p=body();check_rate('setup');expected=app.config.get('ADMIN_SETUP_TOKEN') or ''
  if len(expected)<24 or not hmac.compare_digest(str(p.get('setupToken','')),expected):record_failure('setup');fail('Mã thiết lập không đúng.',403)
  m={'id':'school-owner','username':username(p.get('username')),'password_hash':password(p.get('password')),'name':text(p.get('name')),'email':None,'role':'admin','disabled':0};g.db.execute(insert(members).values(**m));return sign_in(m)
 @app.post('/api/login')
 def login():
  p=body();u=str(p.get('username','')).strip().lower()[:200];check_rate(u);m=one(members,members.c.username==u)
  supplied=p.get('password','');valid=isinstance(supplied,str) and len(supplied)<=200
  if not m or m['disabled'] or not valid or not check_password_hash(m['password_hash'],supplied):record_failure(u);fail('Tên đăng nhập hoặc mật khẩu không đúng.',401)
  g.db.execute(delete(limits).where(limits.c.key==hashlib.sha256(('user:'+u).encode()).hexdigest()));return sign_in(m)
 @app.post('/api/logout')
 def logout():
  g.db.execute(delete(sessions).where(sessions.c.token_hash==g.auth['token_hash']));r=jsonify(ok=True);r.delete_cookie(cookie,path='/');return r
 @app.get('/api/school')
 def snapshot():
  require_member();m=g.member;all_es=rows(enrollments);all_cs=rows(courses);cs=[c for c in all_cs if may_teach(m,c) or (m['role']=='student' and any(e['course_id']==c['id'] and e['member_id']==m['id'] for e in all_es))];ids={c['id'] for c in cs};es=[e for e in all_es if e['course_id'] in ids and (m['role']!='student' or e['member_id']==m['id'])]
  all_ms=rows(members);ms=all_ms
  if m['role']!='admin':ms=[x for x in ms if x['id']==m['id'] or any(c['teacher_id']==x['id'] for c in cs) or (m['role']=='teacher' and any(e['member_id']==x['id'] for e in es))]
  ms=[public_member(x) for x in ms]
  if m['role']!='admin':ms=[{k:v for k,v in x.items() if k not in ['email','username'] or x['id']==m['id']} for x in ms]
  all_rs=rows(records);all_subs=rows(submissions);submitted_ids={sub['assignment_id'] for sub in all_subs if sub['member_id']==m['id']};rs=[];removed=[];active_enrollments={(e['course_id'],e['member_id']) for e in all_es}
  for r in sorted(all_rs,key=lambda x:x['created_at'],reverse=True):
   if r['kind']=='library' or r['course_id'] not in ids:continue
   c=next(c for c in cs if c['id']==r['course_id']);d=json.loads(r['data'])
   if r['kind']=='rosterRemoval':
    if m['role']!='student' and not d.get('restoredAt'):removed.append({k:v for k,v in {**d,'id':r['id'],'course_id':r['course_id']}.items() if k in ['id','course_id','memberId','name','code','removedAt']})
    continue
   if r['kind'] in ['attendance','grade'] and (r['course_id'],d.get('memberId')) not in active_enrollments:continue
   if not may_teach(m,c):
    if r['kind'] in ['attendance','grade','message'] and d.get('memberId') not in [m['id'],'all'] and r['author_id']!=m['id']:continue
    if r['kind']=='outline':
     if not d.get('published'):continue
     d.pop('requirements',None)
    if r['kind']=='assignment':
     d=assigned_exam(d,r['id'],m['id']);d.pop('variants',None)
     show_answers=d.get('allowViewAnswers',False) and r['id'] in submitted_ids
     if not show_answers:
      for field in ['answerNotes','answerGuide','resources']:d.pop(field,None)
     for field in ['questions','essays']:d[field]=[{k:v for k,v in q.items() if k!='outcome' and (k!='correct' or show_answers)} for q in d.get(field,[])]
     if d.get('opensAt') and time.time()<parse_date(d['opensAt']):d.update(questions=[],essays=[]);d.pop('file',None)
   rs.append({**r,'data':d})
  ss=[]
  for sub in all_subs:
   a=next((r for r in all_rs if r['id']==sub['assignment_id']),None)
   if a and a['course_id'] in ids and (a['course_id'],sub['member_id']) in active_enrollments and (m['role']!='student' or sub['member_id']==m['id']):
    sd=with_best_result(json.loads(sub['data']),sub['created_at'])
    if m['role']=='student' and not json.loads(a['data']).get('allowViewScore',True):
     sd['scoreHidden']=True
     for row in [sd]+sd.get('history',[])+sd.get('attemptScores',[])+([sd['bestAttempt']] if sd.get('bestAttempt') else []):
      row['score']=None
      for key in ['quizScore','feedback','gradedAt','gradedBy']:row.pop(key,None)
     sd.update(bestScore=None,bestAttemptCount=None,bestSubmittedAt=None)
    ss.append({**sub,'data':sd})
  lib=[]
  for r in sorted(all_rs,key=lambda x:x['created_at'],reverse=True):
   if r['kind'] not in ['library','outline']:continue
   d=json.loads(r['data'])
   if r['kind']=='outline':
    if r['course_id'] not in ids or not d.get('inLibrary',True):continue
    d['category']='outline'
    if m['role']=='student':d.pop('requirements',None)
   if m['role']=='student' and (r['course_id'] not in ids or not d.get('published')):continue
   c=next((c for c in all_cs if c['id']==r['course_id']),None)
   lib.append({**r,'data':d,'course_label':(c['name']+' · '+c['subject']) if c else 'Lớp / môn đã xóa','author_name':next((x['name'] for x in all_ms if x['id']==r['author_id']),'Giáo viên')})
  return jsonify(me=public_member(dict(m)),removedStudents=removed,members=ms,courses=cs,enrollments=es,records=rs,submissions=ss,library=lib,rosters=[{**r,'data':json.loads(r['data'])} for r in rows(rosters) if m['role']=='admin' or (m['role']=='teacher' and r['author_id']==m['id'] and r['course_id'] in ids)])
 @app.post('/api/student-roster/preview')
 def preview_roster():
  teach(text(request.form.get('courseId')))
  students,errors=parse_roster(request.files.get('file'))
  return jsonify(students=students,errors=errors)
 @app.get('/api/student-roster/template')
 def roster_template():
  require_member()
  if g.member['role'] not in ['admin','teacher']:fail('Chỉ giáo viên và quản trị viên.',403)
  from openpyxl import Workbook
  w=Workbook();sheet=w.active;sheet.title='Danh sach hoc sinh'
  sheet.append(['STT','Mã định danh Bộ GD&ĐT','Họ và tên'])
  sheet.append([1,'HS001','Nguyễn Văn A']);sheet.cell(2,2).number_format='@'
  sheet.column_dimensions['B'].width=30;sheet.column_dimensions['C'].width=35
  out=io.BytesIO();w.save(out);out.seek(0)
  return send_file(out,as_attachment=True,download_name='mau-danh-sach-hoc-sinh.xlsx')
 @app.post('/api/school')
 @app.post('/api/exams')
 @app.post('/api/submissions')
 def action():
  if request.path in ['/api/exams','/api/submissions']:
   try:p=json.loads(request.form.get('payload','{}'))
   except Exception:fail('Dữ liệu đề không hợp lệ.')
   if not isinstance(p,dict):fail('Dữ liệu chưa hợp lệ.')
   if request.path=='/api/exams' and (p.get('action')!='record' or p.get('kind')!='assignment'):fail('Yêu cầu tạo đề chưa hợp lệ.')
   if request.path=='/api/submissions' and p.get('action')!='submit':fail('Yêu cầu nộp bài chưa hợp lệ.')
  else:p=body()
  a=p.get('action');m=g.member
  if a=='submitRoster':
   cid=text(p.get('courseId'));teach(cid);students=roster_students(p.get('students'))
   g.db.execute(insert(rosters).values(id=uid(),course_id=cid,author_id=m['id'],data=json.dumps(students,ensure_ascii=False),status='pending',created_at=now()))
  elif a in ['approveRoster','rejectRoster']:
   admin();rid=text(p.get('id'));r=one(rosters,rosters.c.id==rid)
   if not r or r['status']!='pending':fail('Danh sách đã được xử lý hoặc không tồn tại.',409)
   get_course(r['course_id']);g.db.execute(select(courses).where(courses.c.id==r['course_id']).with_for_update()).first()
   r=g.db.execute(select(rosters).where(rosters.c.id==rid).with_for_update()).mappings().first()
   if not r or r['status']!='pending':fail('Danh sách đã được xử lý hoặc không tồn tại.',409)
   credentials=[]
   if a=='approveRoster':
    for st in roster_students(json.loads(r['data'])):
     existing=one(members,members.c.username==st['code'])
     if existing:
      if existing['role']!='student' or existing['disabled'] or normalized(existing['name'])!=normalized(st['name']):fail(f"Mã {st['code']} đã thuộc tài khoản khác hoặc bị khóa. Hãy kiểm tra trước khi duyệt.",409)
      mid=existing['id'];pw='';note='Tài khoản đã có: giữ mật khẩu hiện tại'
     else:
      mid=uid();pw=secrets.token_urlsafe(12);note='Tài khoản mới'
      g.db.execute(insert(members).values(id=mid,username=st['code'],password_hash=password(pw),name=st['name'],email='',role='student',disabled=0))
     where=and_(enrollments.c.course_id==r['course_id'],enrollments.c.member_id==mid)
     if not one(enrollments,where):g.db.execute(insert(enrollments).values(course_id=r['course_id'],member_id=mid))
     reactivate_removed_student(r['course_id'],mid)
     credentials.append({'name':st['name'],'username':st['code'],'password':pw,'note':note})
   g.db.execute(update(rosters).where(rosters.c.id==rid).values(status='approved' if a=='approveRoster' else 'rejected'))
   return jsonify(ok=True,credentials=credentials)
  elif a=='member':
   admin();role=p.get('role')
   if role not in ['teacher','student']:fail('Vai trò chưa hợp lệ.')
   g.db.execute(insert(members).values(id=uid(),username=username(p.get('username')),password_hash=password(p.get('password')),name=text(p.get('name')),email=text(p.get('email',''),200,False),role=role,disabled=0))
  elif a=='toggleMember':
   admin();x=one(members,members.c.id==text(p.get('id')))
   if not x or x['role']=='admin':fail('Không thể khóa quản trị viên.')
   g.db.execute(update(members).where(members.c.id==x['id']).values(disabled=0 if x['disabled'] else 1));g.db.execute(delete(sessions).where(sessions.c.member_id==x['id']))
  elif a in ['password','resetPassword']:
   if a=='resetPassword':
    admin();target=text(p.get('id'));x=one(members,members.c.id==target)
    if not x or x['role']=='admin':fail('Hãy dùng mục đổi mật khẩu của tôi cho quản trị viên.')
   else:
    target=m['id']
    if not isinstance(p.get('oldPassword'),str) or not check_password_hash(m['password_hash'],p['oldPassword']):fail('Mật khẩu hiện tại chưa đúng.')
   g.db.execute(update(members).where(members.c.id==target).values(password_hash=password(p.get('password'))));g.db.execute(delete(sessions).where(sessions.c.member_id==target))
  elif a=='course':
   admin();tid=text(p.get('teacherId'));t=one(members,and_(members.c.id==tid,members.c.role=='teacher',members.c.disabled==0))
   if not t:fail('Chọn giáo viên đang hoạt động.')
   values={'name':text(p.get('name')),'subject':text(p.get('subject')),'teacher_id':tid}
   if p.get('id'):get_course(p['id']);g.db.execute(update(courses).where(courses.c.id==p['id']).values(**values))
   else:g.db.execute(insert(courses).values(id=uid(),**values))
  elif a=='enroll':
   admin();cid=text(p.get('courseId'));get_course(cid);mid=text(p.get('memberId'));x=one(members,and_(members.c.id==mid,members.c.role=='student',members.c.disabled==0))
   if not x:fail('Chọn học sinh đang hoạt động.')
   g.db.execute(select(courses).where(courses.c.id==cid).with_for_update()).first()
   where=and_(enrollments.c.course_id==cid,enrollments.c.member_id==mid)
   if p.get('remove'):remove_roster_student(cid,mid)
   else:
    if not one(enrollments,where):g.db.execute(insert(enrollments).values(course_id=cid,member_id=mid))
    reactivate_removed_student(cid,mid)
  elif a=='attendanceBatch':
   cid=text(p.get('courseId'));teach(cid);g.db.execute(select(courses).where(courses.c.id==cid).with_for_update()).first()
   day=text(p.get('date'));session=p.get('session');entries=p.get('entries')
   try:
    if datetime.strptime(day,'%Y-%m-%d').strftime('%Y-%m-%d')!=day:raise ValueError()
   except ValueError:fail('Ngày điểm danh chưa hợp lệ.')
   if session not in ['Sáng','Chiều','Tối'] or not isinstance(entries,list):fail('Buổi hoặc danh sách chưa hợp lệ.')
   enrolled_ids={row['member_id'] for row in rows(enrollments) if row['course_id']==cid}
   if not entries or any(not isinstance(row,dict) for row in entries) or len(entries)!=len(enrolled_ids) or {row.get('memberId') for row in entries}!=enrolled_ids:fail('Danh sách lớp đã thay đổi. Tải lại trước khi điểm danh.',409)
   if any(row.get('status') not in ['Có mặt','Vắng có phép','Vắng không phép','Đi muộn','Học bù'] for row in entries):fail('Chọn trạng thái cho đủ học sinh.')
   existing=[r for r in rows(records) if r['kind']=='attendance' and r['course_id']==cid and json.loads(r['data']).get('date')==day and json.loads(r['data']).get('session','Sáng')==session and json.loads(r['data']).get('memberId') in enrolled_ids]
   versions=sorted([{'id':r['id'],'updatedAt':json.loads(r['data']).get('updatedAt','')} for r in existing],key=lambda r:r['id'])
   expected=p.get('expected',[])
   if not isinstance(expected,list) or expected!=versions:fail('Điểm danh đã thay đổi ở phiên khác. Tải lại ngày / buổi rồi điểm danh lại.',409)
   if existing and p.get('replace') is not True:fail('Buổi này đã lưu. Hãy chọn Điểm danh lại.',409)
   stamp=now()+':'+uid()
   for entry in entries:
    d={'memberId':entry['memberId'],'date':day,'session':session,'status':entry['status'],'updatedAt':stamp}
    matches=[r for r in existing if json.loads(r['data']).get('memberId')==entry['memberId']]
    if matches:
     g.db.execute(update(records).where(records.c.id==matches[0]['id']).values(data=json.dumps(d,ensure_ascii=False)))
     for duplicate in matches[1:]:g.db.execute(delete(records).where(records.c.id==duplicate['id']))
    else:g.db.execute(insert(records).values(id=uid(),course_id=cid,kind='attendance',author_id=m['id'],data=json.dumps(d,ensure_ascii=False),created_at=now()))
  elif a=='assignmentSettings':
   r=one(records,and_(records.c.id==text(p.get('id')),records.c.kind=='assignment'))
   if not r:fail('Không tìm thấy đề.',404)
   teach(r['course_id']);g.db.execute(select(courses).where(courses.c.id==r['course_id']).with_for_update()).first()
   r=g.db.execute(select(records).where(records.c.id==r['id']).with_for_update()).mappings().first();d=json.loads(r['data'])
   for key in ['allowViewScore','allowViewAnswers']:
    if key in p:
     if type(p[key]) is not bool:fail('Cấu hình công bố chưa hợp lệ.')
     d[key]=p[key]
   if 'answerGuide' in p:d['answerGuide']=text(p['answerGuide'],30000,False)
   if 'codePolicy' in p:
    policy=p['codePolicy'];fixed=p.get('fixedExamCode','')
    if policy not in ['random','fixed'] or (policy=='fixed' and fixed not in [v['code'] for v in d.get('variants',[])]):fail('Chọn mã đề hợp lệ.')
    if (policy!=d.get('codePolicy','random') or (policy=='fixed' and fixed!=d.get('fixedExamCode'))) and one(submissions,submissions.c.assignment_id==r['id']):fail('Đã có học sinh nộp bài. Không thể đổi cách gán mã đề.',409)
    d.update(codePolicy=policy,fixedExamCode=fixed if policy=='fixed' else '')
   g.db.execute(update(records).where(records.c.id==r['id']).values(data=json.dumps(d,ensure_ascii=False)))
  elif a=='record':
   cid=text(p.get('courseId'));kind=p.get('kind');d=p.get('data')
   if kind not in ['assignment','lesson','schedule','attendance','grade','message'] or not isinstance(d,dict):fail('Dữ liệu chưa hợp lệ.')
   if kind=='message' and m['role']=='student':read(cid);d['memberId']=m['id'];d['read']=False
   else:teach(cid)
   if kind in ['assignment','lesson','schedule','grade','message']:d['title']=text(d.get('title'))
   if kind in ['lesson','assignment','message']:d['body']=text(d.get('body'),20000 if kind!='message' else 4000)
   if d.get('url') and not re.match(r'^https?://',str(d['url']),re.I):fail('Liên kết không hợp lệ.')
   if kind=='assignment':
    validate_exam_data(d);d['subject']=get_course(cid)['subject']
    for key in ['closedAt','closedBy','resources','answerGuide','codePolicy','fixedExamCode']:d.pop(key,None)
    for key,default in [('allowViewScore',True),('allowViewAnswers',False)]:
     if type(d.get(key,default)) is not bool:fail('Cấu hình công bố chưa hợp lệ.')
   if kind=='schedule':parse_date(d.get('at'));d['room']=text(d.get('room'))
   if kind in ['grade','attendance'] or (kind=='message' and m['role']!='student' and d.get('memberId')!='all'):
    mid=text(d.get('memberId'))
    if not one(enrollments,and_(enrollments.c.course_id==cid,enrollments.c.member_id==mid)):fail('Học sinh không thuộc lớp / môn này.')
   if kind=='grade':
    if type(d.get('score')) not in [int,float] or not math.isfinite(d['score']) or not 0<=d['score']<=10:fail('Điểm từ 0 đến 10.')
    if d.get('date'):
     try:
      if datetime.strptime(d['date'],'%Y-%m-%d').strftime('%Y-%m-%d')!=d['date']:raise ValueError()
     except (ValueError,TypeError):fail('Ngày nhập điểm chưa hợp lệ.')
   if kind=='attendance':
    g.db.execute(select(courses).where(courses.c.id==cid).with_for_update()).first()
    d['updatedAt']=now()+':'+uid()
    d['session']=d.get('session','Sáng')
    if d['session'] not in ['Sáng','Chiều','Tối']:fail('Buổi chưa hợp lệ.')
    try:
     if datetime.strptime(d.get('date',''),'%Y-%m-%d').strftime('%Y-%m-%d')!=d.get('date'):raise ValueError()
    except (ValueError,TypeError):fail('Ngày điểm danh chưa hợp lệ.')
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}',str(d.get('date'))) or d.get('status') not in ['Có mặt','Vắng có phép','Vắng không phép','Đi muộn','Học bù']:fail('Điểm danh chưa hợp lệ.')
    prev=next((r for r in rows(records) if r['course_id']==cid and r['kind']=='attendance' and json.loads(r['data']).get('memberId')==d['memberId'] and json.loads(r['data']).get('date')==d['date'] and json.loads(r['data']).get('session','Sáng')==d['session']),None)
    if prev:g.db.execute(update(records).where(records.c.id==prev['id']).values(data=json.dumps(d,ensure_ascii=False)));return jsonify(ok=True)
   if kind=='message':d['read']=False
   rid=uid();file_data=None
   if request.path=='/api/exams':
    file=request.files.get('file')
    if file and file.filename and not(d.get('source')=='uploaded' and d.get('variants')):
     uploaded=exam_upload(file)
     if engine.dialect.name=='postgresql':g.db.execute(sql_text('SELECT pg_advisory_xact_lock(719246831)'))
     used=stored_file_bytes()
     if used+uploaded['size']>250*1024*1024:fail('Kho tệp đã đạt 250 MB.')
     file_data=dict(id=uid(),record_id=rid,**{k:uploaded[k] for k in ['name','mime','size','content']})
     d['file']={'name':uploaded['name'],'mime':uploaded['mime'],'size':uploaded['size'],'url':'/api/exams/'+rid+'/file','preview':uploaded['preview']}
   variant_uploads=[]
   if request.path=='/api/exams' and d.get('source')=='uploaded' and d.get('variants'):
    for v in d['variants']:
     uploaded=exam_upload(request.files.get('file_'+v['code']))
     variant_uploads.append(dict(id=uid(),record_id=rid,code=v['code'],**{k:uploaded[k] for k in ['name','mime','size','content']}))
     v['file']={k:uploaded[k] for k in ['name','mime','size','preview']};v['file']['url']='/api/exams/'+rid+'/files/'+v['code']
    if sum(f['size'] for f in variant_uploads)>25*1024*1024:fail('Tổng tệp mã đề tối đa 25 MB.')
    if engine.dialect.name=='postgresql':g.db.execute(sql_text('SELECT pg_advisory_xact_lock(719246831)'))
    if stored_file_bytes()+sum(f['size'] for f in variant_uploads)>250*1024*1024:fail('Kho tệp đã đạt 250 MB.')
    d['file']=d['variants'][0]['file']
   g.db.execute(insert(records).values(id=rid,course_id=cid,kind=kind,author_id=m['id'],data=json.dumps(d,ensure_ascii=False),created_at=now()))
   for values in variant_uploads:g.db.execute(insert(exam_files).values(**values))
   if file_data:g.db.execute(insert(library_files).values(**file_data))
  elif a=='submit':
   if m['role']!='student':fail('Chỉ học sinh được nộp bài.',403)
   r=one(records,and_(records.c.id==text(p.get('assignmentId')),records.c.kind=='assignment'))
   if not r:fail('Không tìm thấy bài tập.',404)
   read(r['course_id']);g.db.execute(select(courses).where(courses.c.id==r['course_id']).with_for_update()).first()
   r=g.db.execute(select(records).where(records.c.id==r['id']).with_for_update()).mappings().first()
   read(r['course_id']);d=assigned_exam(json.loads(r['data']),r['id'],m['id'])
   if p.get('examCode') is not None and p['examCode']!=d.get('examCode',''):fail('Mã đề đã thay đổi. Mở lại đề trước khi làm bài.',409)
   if d.get('closedAt'):fail('Đề đã đóng, không nhận thêm bài nộp.',409)
   if d.get('opensAt') and time.time()<parse_date(d['opensAt']):fail('Chưa đến giờ làm bài.')
   previous=g.db.execute(select(submissions).where(and_(submissions.c.assignment_id==r['id'],submissions.c.member_id==m['id'])).with_for_update()).mappings().first()
   prev=json.loads(previous['data']) if previous else None
   if previous and d.get('maxAttempts',1)!=0 and prev.get('attemptCount',1)>=d.get('maxAttempts',1):fail('Đã hết số lần làm bài.',409)
   if d.get('due') and time.time()>parse_date(d['due']):fail('Đã hết hạn nộp.')
   uploads=request.files.getlist('files') if request.path=='/api/submissions' else []
   if len(uploads)>5:fail('Tối đa 5 file cho một lần nộp.')
   if uploads and d['type']=='quiz':fail('Đề trắc nghiệm không nhận file tự luận.')
   try:uploads=[read_attachment(file) for file in uploads]
   except ValueError as error:fail(str(error))
   if uploads:
    if engine.dialect.name=='postgresql':g.db.execute(sql_text('SELECT pg_advisory_xact_lock(719246831)'))
    if stored_file_bytes()+sum(file['size'] for file in uploads)>250*1024*1024:fail('Kho tệp đã đạt 250 MB.')
   essay_text=text(p.get('text',''),30000,False)
   if d['type']!='quiz' and not essay_text and not uploads:fail('Cần nhập bài tự luận hoặc đính kèm bài làm.')
   if d['type'] in ['quiz','mixed']:
    ans=p.get('answers');qs=d['questions']
    if not isinstance(ans,list) or len(ans)!=len(qs):fail('Cần trả lời đầy đủ các câu hỏi.')
    earned=sum(auto_question_score(q,ans[i],d.get('tfScoring','progressive'),bool(d.get('composed'))) for i,q in enumerate(qs))
    score=round(earned if d.get('composed') else earned/len(qs)*10,2)
    data={'answers':ans,'score':score,'feedback':''}
    if d['type']=='mixed':data.update(text=essay_text,quizScore=data['score'],score=None)
   else:data={'text':essay_text,'score':None,'feedback':''}
   if d.get('examCode'):data['examCode']=d['examCode']
   data['attemptCount']=(prev.get('attemptCount',1) if prev else 0)+1
   sid=previous['id'] if previous else uid();file_values=[];data['attachments']=[]
   for file in uploads:
    fid=uid();file_values.append(dict(id=fid,submission_id=sid,attempt_count=data['attemptCount'],**file))
    data['attachments'].append({**{k:file[k] for k in ['name','mime','size']},'url':'/api/submissions/'+sid+'/files/'+fid})
   if previous:
    history=prev.get('history',[])
    history.append({k:v for k,v in prev.items() if k not in ['history','attemptScores','bestAttempt','bestScore','bestAttemptCount','bestSubmittedAt']}|{'submittedAt':previous['created_at']})
    data['history']=history[-20:]
    data['attemptScores']=with_best_result(prev,previous['created_at'])['attemptScores']
    data['bestAttempt']=prev.get('bestAttempt')
    with_best_result(data,now())
    g.db.execute(update(submissions).where(submissions.c.id==previous['id']).values(data=json.dumps(data,ensure_ascii=False),created_at=now()))
   else:
    with_best_result(data,now());g.db.execute(insert(submissions).values(id=sid,assignment_id=r['id'],member_id=m['id'],data=json.dumps(data,ensure_ascii=False),created_at=now()))
   for file in file_values:g.db.execute(insert(submission_files).values(**file))
  elif a=='grade':
   sub=one(submissions,submissions.c.id==text(p.get('id')))
   if not sub:fail('Không tìm thấy bài nộp.',404)
   r=one(records,records.c.id==sub['assignment_id']);teach(r['course_id']);score=p.get('score')
   if type(score) not in [int,float] or not math.isfinite(score) or not 0<=score<=10:fail('Điểm từ 0 đến 10.')
   g.db.execute(select(courses).where(courses.c.id==r['course_id']).with_for_update()).first()
   if not one(enrollments,and_(enrollments.c.course_id==r['course_id'],enrollments.c.member_id==sub['member_id'])):fail('Học sinh đã được xóa khỏi lớp. Khôi phục trước khi chấm.',409)
   sub=g.db.execute(select(submissions).where(submissions.c.id==sub['id']).with_for_update()).mappings().first();d=json.loads(sub['data'])
   try:d=grade_attempt(d,sub['created_at'],p.get('attemptCount',d.get('attemptCount',1)),score,text(p.get('feedback',''),4000,False),now(),m['id'])
   except ValueError as error:fail(str(error))
   g.db.execute(update(submissions).where(submissions.c.id==sub['id']).values(data=json.dumps(d,ensure_ascii=False)))
  elif a=='closeAssignment':
   r=one(records,and_(records.c.id==text(p.get('id')),records.c.kind=='assignment'))
   if not r:fail('Không tìm thấy đề.',404)
   teach(r['course_id']);g.db.execute(select(courses).where(courses.c.id==r['course_id']).with_for_update()).first()
   r=g.db.execute(select(records).where(records.c.id==r['id']).with_for_update()).mappings().first();d=json.loads(r['data'])
   if not d.get('closedAt'):
    d.update(closedAt=now(),closedBy=m['id']);g.db.execute(update(records).where(records.c.id==r['id']).values(data=json.dumps(d,ensure_ascii=False)))
  elif a=='removeRosterStudent':remove_roster_student(text(p.get('courseId')),text(p.get('memberId')))
  elif a=='restoreRosterStudent':restore_roster_student(text(p.get('id')))
  elif a=='readMessage':
   r=one(records,and_(records.c.id==text(p.get('id')),records.c.kind=='message'))
   if not r:fail('Không tìm thấy trao đổi.')
   c=read(r['course_id']);d=json.loads(r['data'])
   if may_teach(m,c) and r['author_id']!=m['id']:d['read']=True
   elif m['role']=='student' and d.get('memberId') in [m['id'],'all'] and r['author_id']!=m['id']:d['readBy']=list(set(d.get('readBy',[])+[m['id']]))
   else:fail('Bạn không phải người nhận.',403)
   g.db.execute(update(records).where(records.c.id==r['id']).values(data=json.dumps(d,ensure_ascii=False)))
  elif a=='deleteRecord':
   r=one(records,records.c.id==text(p.get('id')))
   if not r:fail('Không tìm thấy nội dung.')
   if r['kind']=='library':
    if m['role']!='admin' and (m['role']!='teacher' or r['author_id']!=m['id']):fail('Chỉ người đăng hoặc quản trị viên được xóa tài liệu.',403)
    g.db.execute(delete(library_files).where(library_files.c.record_id==r['id']))
   else:
    teach(r['course_id']);g.db.execute(select(courses).where(courses.c.id==r['course_id']).with_for_update()).first()
   g.db.execute(delete(library_files).where(library_files.c.record_id==r['id']))
   g.db.execute(delete(exam_files).where(exam_files.c.record_id==r['id']))
   g.db.execute(delete(submission_files).where(submission_files.c.submission_id.in_(select(submissions.c.id).where(submissions.c.assignment_id==r['id']))))
   g.db.execute(delete(exam_resources).where(exam_resources.c.record_id==r['id']))
   g.db.execute(delete(submissions).where(submissions.c.assignment_id==r['id']));g.db.execute(delete(records).where(records.c.id==r['id']))
  else:fail('Thao tác chưa hợp lệ.')
  return jsonify(ok=True)
 @app.post('/api/library')
 def save_library():
  if g.member['role'] not in ['admin','teacher']:fail('Chỉ giáo viên và quản trị viên được đăng tài liệu.',403)
  f=request.form;rid=f.get('id') or uid();existing=one(records,records.c.id==rid)
  if existing and (existing['kind']!='library' or (g.member['role']!='admin' and existing['author_id']!=g.member['id'])):fail('Bạn không được sửa tài liệu này.',403)
  cid=text(f.get('courseId'));c=teach(cid);category=f.get('category');grade=f.get('grade','')
  if category not in ['lesson','exam'] or grade not in [str(i) for i in range(1,13)]+['other']:fail('Chọn loại tài liệu và khối lớp hợp lệ.')
  d={'title':text(f.get('title')), 'body':text(f.get('body',''),10000,False),'url':text(f.get('url',''),2000,False),'category':category,'grade':grade,'subject':c['subject'],'published':f.get('published')=='true'}
  if d['url'] and not re.match(r'^https?://',d['url'],re.I):fail('Liên kết cần bắt đầu bằng http:// hoặc https://.')
  if existing:d['file']=json.loads(existing['data']).get('file')
  file=request.files.get('file');file_data=None
  if file and file.filename:
   name=secure_filename(file.filename) or 'tai-lieu';ext=Path(name).suffix.lower();allowed={'.pdf':'application/pdf','.docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','.xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','.txt':'text/plain','.md':'text/markdown','.png':'image/png','.jpg':'image/jpeg','.jpeg':'image/jpeg'}
   if ext not in allowed:fail('Hỗ trợ PDF, DOCX, XLSX, TXT, MD và ảnh PNG/JPG.')
   content=file.stream.read(5*1024*1024+1)
   if not content or len(content)>5*1024*1024:fail('Tệp cần có nội dung và không vượt quá 5 MB.')
   file_data={'id':uid(),'record_id':rid,'name':name,'mime':allowed[ext],'size':len(content),'content':content}
   if engine.dialect.name=='postgresql':g.db.execute(sql_text('SELECT pg_advisory_xact_lock(719246831)'))
   used=stored_file_bytes();old_size=g.db.execute(select(library_files.c.size).where(library_files.c.record_id==rid)).scalar_one_or_none() or 0
   if used-old_size+len(content)>250*1024*1024:fail('Thư viện đã đạt hạn mức tệp 250 MB. Hãy xóa tệp cũ hoặc dùng liên kết.')
   d['file']={'name':name,'size':len(content),'url':'/api/library/'+rid+'/file'}
  elif f.get('removeFile')=='true':d['file']=None
  if not d['body'] and not d['url'] and not d.get('file'):fail('Cần nội dung, liên kết hoặc tệp tài liệu.')
  values={'course_id':cid,'data':json.dumps(d,ensure_ascii=False)}
  if existing:g.db.execute(update(records).where(records.c.id==rid).values(**values))
  else:g.db.execute(insert(records).values(id=rid,kind='library',author_id=g.member['id'],created_at=now(),**values))
  if file_data or f.get('removeFile')=='true':g.db.execute(delete(library_files).where(library_files.c.record_id==rid))
  if file_data:g.db.execute(insert(library_files).values(**file_data))
  return jsonify(ok=True,id=rid)
 @app.get('/api/library/<rid>/file')
 def download_library(rid):
  require_member();r=one(records,and_(records.c.id==rid,records.c.kind=='library'))
  if not r:fail('Không tìm thấy tài liệu.',404)
  d=json.loads(r['data'])
  if g.member['role']=='student':
   if not d.get('published'):fail('Tài liệu chưa được xuất bản.',403)
   read(r['course_id'])
  f=one(library_files,library_files.c.record_id==rid)
  if not f:fail('Không tìm thấy tệp.',404)
  return send_file(io.BytesIO(f['content']),mimetype=f['mime'],download_name=f['name'],as_attachment=True,max_age=0)
 @app.get('/api/submissions/<sid>/files/<fid>')
 def submission_attachment(sid,fid):
  require_member();sub=one(submissions,submissions.c.id==sid)
  if not sub:fail('Không tìm thấy bài nộp.',404)
  r=one(records,records.c.id==sub['assignment_id']);c=read(r['course_id'])
  if not may_teach(g.member,c) and sub['member_id']!=g.member['id']:fail('Bạn không được xem bài nộp này.',403)
  if not one(enrollments,and_(enrollments.c.course_id==r['course_id'],enrollments.c.member_id==sub['member_id'])):fail('Học sinh đã rút khỏi lớp.',403)
  file=one(submission_files,and_(submission_files.c.id==fid,submission_files.c.submission_id==sid))
  if not file:fail('Không tìm thấy file.',404)
  return send_file(io.BytesIO(file['content']),mimetype=file['mime'],download_name=file['name'],as_attachment=not(request.args.get('preview')=='1' and (file['mime']=='application/pdf' or file['mime'].startswith('image/'))),max_age=0)
 @app.post('/api/exams/<rid>/resources')
 def upload_exam_resource(rid):
  r=one(records,and_(records.c.id==rid,records.c.kind=='assignment'))
  if not r:fail('Không tìm thấy đề.',404)
  teach(r['course_id']);kind=request.form.get('kind')
  if kind not in ['answer','solution']:fail('Loại tài liệu chưa hợp lệ.')
  source=request.files.get('file')
  if source and Path(source.filename or '').suffix.lower()=='.docx':file=exam_upload(source)
  else:
   try:file=read_attachment(source)
   except ValueError as error:fail(str(error))
  g.db.execute(select(courses).where(courses.c.id==r['course_id']).with_for_update()).first()
  r=g.db.execute(select(records).where(records.c.id==rid).with_for_update()).mappings().first();d=json.loads(r['data'])
  if engine.dialect.name=='postgresql':g.db.execute(sql_text('SELECT pg_advisory_xact_lock(719246831)'))
  old=one(exam_resources,and_(exam_resources.c.record_id==rid,exam_resources.c.kind==kind))
  if stored_file_bytes()-(old['size'] if old else 0)+file['size']>250*1024*1024:fail('Kho tệp đã đạt 250 MB.')
  g.db.execute(delete(exam_resources).where(and_(exam_resources.c.record_id==rid,exam_resources.c.kind==kind)))
  g.db.execute(insert(exam_resources).values(id=uid(),record_id=rid,kind=kind,**{k:file[k] for k in ['name','mime','size','content']}))
  d.setdefault('resources',{})[kind]={**{k:file[k] for k in ['name','mime','size']},'url':'/api/exams/'+rid+'/resources/'+kind}
  g.db.execute(update(records).where(records.c.id==rid).values(data=json.dumps(d,ensure_ascii=False)));return jsonify(ok=True)
 @app.get('/api/exams/<rid>/resources/<kind>')
 def download_exam_resource(rid,kind):
  require_member();r=one(records,and_(records.c.id==rid,records.c.kind=='assignment'))
  if not r:fail('Không tìm thấy đề.',404)
  c=read(r['course_id']);d=json.loads(r['data'])
  if not may_teach(g.member,c) and (not d.get('allowViewAnswers',False) or not one(submissions,and_(submissions.c.assignment_id==rid,submissions.c.member_id==g.member['id']))):fail('Đáp án chưa được công bố cho bạn.',403)
  file=one(exam_resources,and_(exam_resources.c.record_id==rid,exam_resources.c.kind==kind))
  if not file:fail('Không tìm thấy tài liệu.',404)
  return send_file(io.BytesIO(file['content']),mimetype=file['mime'],download_name=file['name'],as_attachment=True,max_age=0)
 @app.get('/api/exams/<rid>/files/<code>')
 def variant_file(rid,code):
  require_member();r=one(records,and_(records.c.id==rid,records.c.kind=='assignment'))
  if not r:fail('Không tìm thấy đề.',404)
  c=read(r['course_id']);d=json.loads(r['data'])
  if not may_teach(g.member,c) and d.get('opensAt') and time.time()<parse_date(d['opensAt']):fail('Chưa đến giờ mở đề.',403)
  if not may_teach(g.member,c) and assigned_exam(d,rid,g.member['id']).get('examCode')!=code:fail('Bạn không được tải mã đề khác.',403)
  f=one(exam_files,and_(exam_files.c.record_id==rid,exam_files.c.code==code))
  if not f:fail('Không tìm thấy tệp mã đề.',404)
  return send_file(io.BytesIO(f['content']),mimetype=f['mime'],download_name=f['name'],as_attachment=not(request.args.get('preview')=='1' and f['mime']=='application/pdf'),max_age=0)
 @app.get('/api/exams/key-template')
 def exam_key_template():
  from openpyxl import Workbook
  require_member()
  if g.member['role']=='student':fail('Chỉ giáo viên dùng mẫu đáp án.',403)
  wb=Workbook();ws=wb.active;ws.title='Đáp án TN Maker';ws.append(['Câu\\Mã Đề','101','102','103','104'])
  for count,answer in [(12,'A'),(4,'ĐĐSS'),(6,'0,25')]:
   for number in range(1,count+1):ws.append([number]+[answer]*4)
  ws.freeze_panes='B2'
  for col in 'ABCDE':ws.column_dimensions[col].width=22
  out=io.BytesIO();wb.save(out);out.seek(0)
  return send_file(out,as_attachment=True,download_name='mau-dap-an-tn-maker.xlsx')
 @app.post('/api/exams/keys')
 def exam_keys():
  teach(text(request.form.get('courseId')));file=request.files.get('file')
  if not file or not file.filename.lower().endswith('.xlsx'):fail('Chọn file đáp án TN Maker .xlsx.')
  raw=file.stream.read(2*1024*1024+1)
  if len(raw)>2*1024*1024:fail('File đáp án tối đa 2 MB.')
  try:
   return jsonify(rows=read_answer_keys(raw))
  except ValueError as error:fail(str(error))
  except Exception:fail('Không đọc được Excel đáp án. Hãy dùng file TN Maker .xlsx không có mật khẩu.')
 @app.post('/api/exams/inspect')
 def inspect_exam():
  teach(text(request.form.get('courseId')));uploaded=exam_upload(request.files.get('file'))
  return jsonify(**{k:uploaded[k] for k in ['name','mime','size','preview']})
 @app.post('/api/exams/convert')
 def convert_exam_word():
  teach(text(request.form.get('courseId')));file=request.files.get('file')
  if not file or not file.filename.lower().endswith('.docx'):fail('Chọn đề Word .docx để chuyển sang PDF.')
  uploaded=exam_upload(file)
  response=send_file(io.BytesIO(uploaded['content']),mimetype='application/pdf',download_name=uploaded['name'],as_attachment=True,max_age=0)
  response.headers['Cache-Control']='no-store'
  return response
 @app.post('/api/outlines')
 def save_outline():
  p=body();cid=text(p.get('courseId'));c=teach(cid);d=outline_data(p.get('data') or {});d['subject']=c['subject'];rid=text(p.get('id')) if p.get('id') else uid();existing=one(records,records.c.id==rid)
  if existing:
   if existing['kind']!='outline':fail('Không tìm thấy đề cương.',404)
   teach(existing['course_id'])
   g.db.execute(update(records).where(records.c.id==rid).values(course_id=cid,data=json.dumps(d,ensure_ascii=False)))
  else:g.db.execute(insert(records).values(id=rid,course_id=cid,kind='outline',author_id=g.member['id'],data=json.dumps(d,ensure_ascii=False),created_at=now()))
  return jsonify(ok=True,id=rid)
 @app.post('/api/outlines/export')
 def export_outline_draft():
  from exam_export import standalone_outline_doc
  p=body();c=teach(text(p.get('courseId')));d=outline_data(p.get('data') or {});d['subject']=c['subject']
  return send_file(io.BytesIO(standalone_outline_doc(d)),as_attachment=True,download_name='de-cuong-on-tap.docx')
 @app.get('/api/outlines/<rid>/word')
 def export_saved_outline(rid):
  from exam_export import standalone_outline_doc
  require_member();r=one(records,and_(records.c.id==rid,records.c.kind=='outline'))
  if not r:fail('Không tìm thấy đề cương.',404)
  read(r['course_id']);d=json.loads(r['data'])
  if g.member['role']=='student' and not d.get('published'):fail('Đề cương chưa công bố.',403)
  return send_file(io.BytesIO(standalone_outline_doc(d)),as_attachment=True,download_name='de-cuong-on-tap.docx')
 @app.post('/api/exams/export')
 def export_exam():
  from exam_export import export_pack,outline_doc
  p=body();c=teach(text(p.get('courseId')));d=p.get('data')
  if not isinstance(d,dict) or not d.get('composed'):fail('Cần đề soạn trực tiếp để xuất tệp.')
  d['title']=text(d.get('title'));d['body']=text(d.get('body',''),12000,False)
  validate_exam_data(d);d['subject']=c['subject']
  if p.get('format')=='outline':return send_file(io.BytesIO(outline_doc(d)),mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',as_attachment=True,download_name='de-cuong.docx')
  return send_file(io.BytesIO(export_pack(d)),mimetype='application/zip',as_attachment=True,download_name='bo-de-kiem-tra.zip')
 @app.get('/api/exams/<rid>/file')
 def exam_file(rid):
  require_member();r=one(records,and_(records.c.id==rid,records.c.kind=='assignment'))
  if not r:fail('Không tìm thấy đề.',404)
  c=read(r['course_id']);d=json.loads(r['data'])
  if not may_teach(g.member,c) and d.get('opensAt') and time.time()<parse_date(d['opensAt']):fail('Chưa đến giờ mở đề.',403)
  f=one(library_files,library_files.c.record_id==rid)
  if not f:fail('Không tìm thấy tệp đề.',404)
  inline=request.args.get('preview')=='1' and f['mime']=='application/pdf'
  response=send_file(io.BytesIO(f['content']),mimetype=f['mime'],download_name=f['name'],as_attachment=not inline,max_age=0)
  return response
 @app.get('/')
 def index():return send_from_directory(app.static_folder,'index.html')
 @app.get('/assets/<path:name>')
 def assets(name):return send_from_directory(Path(app.static_folder)/'assets',name,max_age=31536000)
 return app
