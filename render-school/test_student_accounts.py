import io,json
from sqlalchemy import select,update,delete
from openpyxl import load_workbook
from test_school import app,client,setup,action,login,state,token,post
from server import members,student_credentials,account_audits,rosters
from student_credentials import seal_password,open_password

def prepare(admin,app):
 setup(admin)
 for user,role in [('teach1','teacher'),('teach2','teacher'),('student1','student'),('student2','student')]:assert action(admin,{'action':'member','username':user,'name':user,'role':role,'password':'Password123!'}).status_code==200
 ids={m['username']:m['id'] for m in state(admin)['members']}
 for name,teacher in [('10A','teach1'),('10B','teach2')]:assert action(admin,{'action':'course','name':name,'subject':'Toán','teacherId':ids[teacher]}).status_code==200
 cs={c['name']:c['id'] for c in state(admin)['courses']}
 for cid,u in [(cs['10A'],'student1'),(cs['10B'],'student2')]:assert action(admin,{'action':'enroll','courseId':cid,'memberId':ids[u]}).status_code==200
 clients=[app.test_client() for _ in range(4)]
 for c,u in zip(clients,['teach1','teach2','student1','student2']):assert login(c,u).status_code==200
 return *clients,ids,cs

def edit(c,mid,old,new,password=''):
 return post(c,'/api/student-accounts/'+mid,{'username':new,'expectedUsername':old,'password':password},token(c))

def test_teacher_scope_reveal_encryption_export_and_audit(client,app):
 t1,t2,s1,s2,ids,cs=prepare(client,app)
 listed=t1.get('/api/student-accounts');assert listed.status_code==200
 assert [s['id'] for s in listed.json['students']]==[ids['student1']]
 assert listed.json['students'][0]['passwordAvailable']
 assert 'Password123!' not in listed.text and 'password_hash' not in listed.text and 'encrypted_password' not in listed.text
 assert post(t1,'/api/student-accounts/'+ids['student1']+'/reveal',{},token(t1)).json['password']=='Password123!'
 assert post(t1,'/api/student-accounts/'+ids['student2']+'/reveal',{},token(t1)).status_code==403
 assert s1.get('/api/student-accounts').status_code==403
 assert post(s1,'/api/student-accounts/'+ids['student1']+'/reveal',{},token(s1)).status_code==403
 assert edit(t1,ids['teach1'],'teach1','bad').status_code==403
 assert edit(t2,ids['student1'],'student1','bad').status_code==403
 assert post(t1,'/api/student-accounts/export',{'ids':[ids['student1'],ids['student2']],'includePasswords':True},token(t1)).status_code==403
 result=post(t1,'/api/student-accounts/export',{'ids':[ids['student1']],'includePasswords':True},token(t1));assert result.status_code==200 and result.headers['Cache-Control']=='no-store'
 wb=load_workbook(io.BytesIO(result.data));ws=wb.active;assert ws['B2'].value=='student1' and ws['D2'].value=='Password123!'
 with app.extensions['school_engine'].connect() as db:
  vault=db.execute(select(student_credentials).where(student_credentials.c.member_id==ids['student1'])).mappings().one()
  assert 'Password123!' not in vault['encrypted_password']
  audits=db.execute(select(account_audits)).mappings().all();assert {a['action'] for a in audits}>={'password_reveal','password_export'}
  assert 'Password123!' not in json.dumps([dict(a) for a in audits])

def test_rename_preserves_identity_scores_rosters_and_login_and_rejects_duplicates(client,app):
 t1,t2,s1,s2,ids,cs=prepare(client,app)
 assert action(t1,{'action':'submitRoster','courseId':cs['10A'],'students':[{'code':'student1','name':'student1'}]}).status_code==200
 assert action(t1,{'action':'record','kind':'grade','courseId':cs['10A'],'data':{'title':'Bài 1','score':9,'memberId':ids['student1']}}).status_code==200
 assert edit(t1,ids['student1'],'student1','student2').status_code==409
 assert edit(t1,ids['student1'],'student1','000123').status_code==200
 assert s1.get('/api/school').status_code==401
 assert login(s1,'student1').status_code==401
 assert login(s1,'000123').status_code==200
 assert state(s1)['me']['id']==ids['student1']
 assert next(r for r in state(s1)['records'] if r['kind']=='grade')['data']['score']==9
 assert state(t1)['rosters'][0]['data'][0]['code']=='000123'
 assert edit(t1,ids['student1'],'student1','another').status_code==409
 exported=post(t1,'/api/student-accounts/export',{'ids':[ids['student1']]},token(t1));wb=load_workbook(io.BytesIO(exported.data));assert wb.active['B2'].value=='000123' and wb.active['B2'].data_type=='s'
 assert edit(t1,ids['student1'],'000123','000123','=SecretPass123').status_code==200
 assert login(s1,'000123','=SecretPass123').status_code==200
 export=post(t1,'/api/student-accounts/export',{'ids':[ids['student1']],'includePasswords':True},token(t1));cell=load_workbook(io.BytesIO(export.data)).active['D2'];assert cell.value=='=SecretPass123' and cell.data_type=='s'

def test_old_hash_is_not_recoverable_and_self_password_change_revokes_staff_view(client,app):
 t1,t2,s1,s2,ids,cs=prepare(client,app)
 with app.extensions['school_engine'].begin() as db:db.execute(delete(student_credentials).where(student_credentials.c.member_id==ids['student1']))
 assert not t1.get('/api/student-accounts').json['students'][0]['passwordAvailable']
 assert post(t1,'/api/student-accounts/'+ids['student1']+'/reveal',{},token(t1)).status_code==409
 assert action(t1,{'action':'resetPassword','id':ids['student1'],'password':'IssuedNew123!'}).status_code==200
 assert post(t1,'/api/student-accounts/'+ids['student1']+'/reveal',{},token(t1)).json['password']=='IssuedNew123!'
 assert login(s1,'student1','IssuedNew123!').status_code==200
 assert action(s1,{'action':'password','oldPassword':'IssuedNew123!','password':'PersonalNew123!'}).status_code==200
 assert post(t1,'/api/student-accounts/'+ids['student1']+'/reveal',{},token(t1)).status_code==409
 assert not t1.get('/api/student-accounts').json['students'][0]['passwordAvailable']
 assert action(t1,{'action':'resetPassword','id':ids['teach2'],'password':'NotAllowed123!'}).status_code==403

def test_cipher_is_bound_to_member_hash_and_secret():
 secret='s'*40;encrypted=seal_password(secret,'one','test-password','HASH')
 assert open_password(secret,'one',encrypted,'HASH')=='test-password'
 assert open_password(secret,'two',encrypted,'HASH') is None
 assert open_password(secret,'one',encrypted,'DIFFERENT') is None
 assert open_password('different'*8,'one',encrypted,'HASH') is None

def test_six_character_issued_password_and_personal_change(client,app):
 t1,t2,s1,s2,ids,cs=prepare(client,app)
 assert edit(t1,ids['student1'],'student1','student1','Ab123').status_code==400
 assert edit(t1,ids['student1'],'student1','student1','Ab1234').status_code==200
 assert login(s1,'student1','Ab1234').status_code==200
 assert post(t1,'/api/student-accounts/'+ids['student1']+'/reveal',{},token(t1)).json['password']=='Ab1234'
 assert action(s1,{'action':'password','oldPassword':'Ab1234','password':'Cd5678'}).status_code==200
 assert login(s1,'student1','Cd5678').status_code==200
 assert login(s2,'student2','Password123!').status_code==200
