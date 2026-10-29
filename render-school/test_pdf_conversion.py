import io
import shutil
import subprocess
import pytest
from docx import Document
import pdf_conversion
from pdf_conversion import convert_docx_to_pdf, ConversionError


def word_bytes():
    doc = Document()
    doc.add_heading('Đề kiểm tra Toán', 0)
    doc.add_paragraph('Câu 1. Giải phương trình x² − 4 = 0. Đáp án A, B, C, D.')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Câu'
    table.cell(0, 1).text = 'Nội dung'
    table.cell(1, 0).text = '2'
    table.cell(1, 1).text = 'Hình học và số học'
    doc.add_page_break()
    doc.add_paragraph('Trang 2. Hết đề.')
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_unavailable_converter_has_actionable_error(monkeypatch):
    monkeypatch.setattr(pdf_conversion, 'converter_path', lambda: None)
    with pytest.raises(ConversionError, match='chưa có bộ chuyển'):
        convert_docx_to_pdf(word_bytes())


def fake_process(monkeypatch, tmp_path, output, returncode=0):
    monkeypatch.setattr(pdf_conversion, 'converter_path', lambda: '/usr/bin/soffice')
    captured = {}
    class Process:
        pid = 999999
        def __init__(self, command, **kwargs):
            captured.update(command=command, **kwargs)
            self.returncode = returncode
            directory = command[command.index('--outdir') + 1]
            if output is not None:
                (pdf_conversion.Path(directory) / 'exam.pdf').write_bytes(output)
        def communicate(self, timeout=None):
            return b'', b''
    monkeypatch.setattr(pdf_conversion.subprocess, 'Popen', Process)
    return captured


@pytest.mark.parametrize('output,returncode', [(None, 1), (b'not PDF', 0), (b'%PDF-1.4\ntruncated', 0), (b'x' * (5*1024*1024+1), 0)])
def test_conversion_failure_never_returns_word(monkeypatch, tmp_path, output, returncode):
    fake_process(monkeypatch, tmp_path, output, returncode)
    with pytest.raises(ConversionError):
        convert_docx_to_pdf(word_bytes())


def test_converter_receives_no_server_secrets_and_cleans_files(monkeypatch, tmp_path):
    monkeypatch.setenv('DATABASE_URL', 'private-connection')
    monkeypatch.setenv('SECRET_KEY', 'private-key')
    captured = fake_process(monkeypatch, tmp_path, b'%PDF-1.4\ncontent\n%%EOF')
    assert convert_docx_to_pdf(word_bytes()).startswith(b'%PDF-')
    assert 'DATABASE_URL' not in captured['env'] and 'SECRET_KEY' not in captured['env']
    assert not pdf_conversion.Path(captured['env']['HOME']).exists()
    assert '--headless' in captured['command']


def test_only_one_conversion_runs_at_a_time(monkeypatch, tmp_path):
    fake_process(monkeypatch, tmp_path, b'%PDF-1.4\ncontent\n%%EOF')
    with open(pdf_conversion.Path(pdf_conversion.tempfile.gettempdir()) / 'school-word-pdf.lock', 'a') as lock:
        pdf_conversion.fcntl.flock(lock, pdf_conversion.fcntl.LOCK_EX | pdf_conversion.fcntl.LOCK_NB)
        with pytest.raises(ConversionError, match='đang chuyển'):
            convert_docx_to_pdf(word_bytes())


def test_timeout_stops_conversion_process_group(monkeypatch, tmp_path):
    fake_process(monkeypatch, tmp_path, None)
    process_type = pdf_conversion.subprocess.Popen
    def communicate(self, timeout=None):
        if timeout:
            raise subprocess.TimeoutExpired('soffice', timeout)
        return b'', b''
    monkeypatch.setattr(process_type, 'communicate', communicate)
    killed = []
    monkeypatch.setattr(pdf_conversion.os, 'killpg', lambda *args: killed.append(args))
    with pytest.raises(ConversionError, match='quá thời gian'):
        convert_docx_to_pdf(word_bytes())
    assert killed == [(999999, pdf_conversion.signal.SIGKILL)]


@pytest.mark.skipif(not (shutil.which('libreoffice') or shutil.which('soffice')), reason='LibreOffice is not installed')
def test_real_word_conversion():
    pdf = convert_docx_to_pdf(word_bytes())
    assert pdf.startswith(b'%PDF-') and b'%%EOF' in pdf[-1024:]
